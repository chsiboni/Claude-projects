#!/usr/bin/env python3
"""
Fizikal — הצלבת לקוחות הנה"ח עם הייצוא הטרי של משה (סניפים + עובדים).

    python3 src/fresh_export.py --month 2026-07 --tag 2026-09

קלט: שני קבצי CSV מהמערכת של Fizikal (cp1255), לפי המפרט ב-out/Fizikal_בקשה_למתכנת.xlsx:
  · סניפים — שורה לכל סניף, עם ח.פ ושם בהנה"ח *של הסניף* ושם החברה הרשמי.
  · עובדים — מנהל ומעלה, פעילים, שורה לכל שיוך עובד-סניף.

פלט:
  · out/sap_to_company_<tag>.json   — לכל מספר סאפ: קודי חברה, סניפים בתחום, איך הוצלב.
  · out/migration_<tag>.json        — לקוחות המעבר (בינוני/קטן/מוב, לא "לא פעיל") עם אנשי הקשר.
  · out/Fizikal_מעבר_<tag>.xlsx     — אותו דבר לעין אנושית.

מה הייצוא שבר, ומה תוקן פה:
  · טלפון, ת.ז וח.פ יצאו כמספרים — האפס המוביל נחתך. ח.פ של 8 ספרות ות.ז של 6-8
    ספרות מרופדים לתשע, ות.ז מאומתת בספרת ביקורת אחרי הריפוד.
  · ת.ז חסרה יצאה כ-'0' ולא כתא ריק. '0' נפסל במפורש (000000000 עובר ספרת ביקורת).
  · טלפונים בשלושה פורמטים: 9 ספרות בלי אפס, 052-1234567, וקווי נייחים. רק נייד תקין נשמר.
  · כל השדות מרופדים ברווחים — הכל נחתך.
"""

import argparse
import csv
import io
import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from link_clients import JUNK_HP, norm, pad, read_families, link
from revenue_pipeline import OUT, ROOT, load_config

DATA = ROOT / "data" / "moshe_2026-09"
MIGRATION_TIERS = ("בינוני", "קטן", "מוב")
RANK = {"הרשאת על": 3, "ניהול חברה": 2, "מנהל": 1, "מנהל משמרת": 0}


# ---------- normalisation ----------

def read_csv(path):
    text = Path(path).read_bytes().decode("cp1255")
    rows = list(csv.reader(io.StringIO(text)))
    header = [h.strip() for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not any(c.strip() for c in r):
            continue
        out.append({h: c.strip() for h, c in zip(header, r) if h})
    return out


def valid_id(t):
    """ספרת ביקורת של ת.ז. '0' ואפסים בלבד נפסלים במפורש — הם עוברים את החישוב."""
    if not t.isdigit() or len(t) != 9 or not t.strip("0"):
        return False
    total = 0
    for i, ch in enumerate(t):
        d = int(ch) * (1 if i % 2 == 0 else 2)
        total += d if d < 10 else d - 9
    return total % 10 == 0


def norm_tz(v):
    t = re.sub(r"\D", "", v or "")
    if 6 <= len(t) <= 8:
        t = t.zfill(9)
    return t if valid_id(t) else ""


def norm_phone(v):
    p = re.sub(r"\D", "", v or "")
    if len(p) == 12 and p.startswith("972"):
        p = "0" + p[3:]
    if len(p) == 9 and p.startswith("5"):
        p = "0" + p
    return p if re.fullmatch(r"05\d{8}", p) else ""


def norm_email(v):
    e = (v or "").strip().lower()
    return e if re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", e) else ""


def norm_hp(v):
    return pad(re.sub(r"\D", "", v or ""))


# ---------- the two exports ----------

def load_branches(path):
    rows = read_csv(path)
    for r in rows:
        r["code"] = re.sub(r"\D", "", r.get("חברה", ""))
        r["branch"] = re.sub(r"\D", "", r.get("סניף", ""))
        r["hp"] = norm_hp(r.get("ח.פ", ""))
        r["active"] = r.get("פעיל", "") == "פעיל"
    return [r for r in rows if r["code"]]


def load_employees(path):
    rows = read_csv(path)
    for r in rows:
        r["code"] = re.sub(r"\D", "", r.get("חברה", ""))
        r["branch"] = re.sub(r"\D", "", r.get("סניף", ""))
        r["emp"] = r.get("מספר עובד", "")
        r["phone"] = norm_phone(r.get("טלפון", ""))
        r["tz"] = norm_tz(r.get("ת.ז.", ""))
        r["email"] = norm_email(r.get('דוא"ל', ""))
        r["role"] = r.get("הרשאה", "")
        r["name"] = re.sub(r"\s+", " ", r.get("שם עובד", "")).strip()
    return [r for r in rows if r["code"]]


def companies_index(branches):
    """
    אותו מבנה כמו link_clients.read_companies, אבל מהייצוא הטרי:
    ח.פ ושם בהנה"ח ושם הסניף → (קוד, סניף); שם החברה הרשמי → (קוד, "") כלומר כל החברה.
    """
    by_hp, by_name = defaultdict(set), defaultdict(set)
    branch_names = {}
    for b in branches:
        key = (b["code"], b["branch"])
        branch_names[key] = b.get("שם הסניף", "")
        if b["hp"]:
            by_hp[b["hp"]].add(key)
        for col in ("שם בהנהח", "שם הסניף"):
            if norm(b.get(col)):
                by_name[norm(b.get(col))].add(key)
        if norm(b.get("שם חברה")):
            by_name[norm(b.get("שם חברה"))].add((b["code"], ""))
    return by_hp, by_name, branch_names


# ---------- people per client ----------

def attach_people(linked, employees, branch_names):
    """
    לכל לקוח מקושר: העובדים של הסניפים שבתחום שלו. תחום ["*"] = כל סניפי החברה.
    אדם שמנהל כמה סניפים של אותו לקוח מופיע פעם אחת, על הסניף הראשון.
    """
    by_key = defaultdict(list)
    for e in employees:
        by_key[(e["code"], e["branch"])].append(e)
    by_code = defaultdict(list)
    for e in employees:
        by_code[e["code"]].append(e)

    for sap, rec in linked.items():
        seen, people = set(), []
        for code, branches in rec["scope"].items():
            pool = by_code[code] if branches == ["*"] else [e for br in branches for e in by_key[(code, br)]]
            for e in sorted(pool, key=lambda e: -RANK.get(e["role"], 0)):
                pid = (e["code"], e["emp"])
                if pid in seen:
                    continue
                seen.add(pid)
                people.append({
                    "name": e["name"] or "ללא שם", "role": e["role"], "emp": e["emp"],
                    "code": e["code"], "branch_no": e["branch"],
                    "branch": branch_names.get((e["code"], e["branch"]), ""),
                    "phone": e["phone"], "email": e["email"], "tz": e["tz"],
                    "ok": bool(e["email"] and e["tz"]),
                })
        rec["people"] = people
        rec["n"] = len(people)
        rec["ready"] = sum(1 for p in people if p["ok"])
        rec["fit"] = "אין אנשי קשר" if not people else ("מוכן" if rec["ready"] else "חסר פרט")
    return linked


def migration_set(clients, linked):
    """לקוחות המעבר: בינוני/קטן/מוב שאינם 'לא פעיל'. מי שלא הוצלב מופיע בלי אנשי קשר."""
    out = []
    for c in clients:
        if c["tier"] not in MIGRATION_TIERS or c["status"] == "לא פעיל":
            continue
        rec = linked.get(c["sap"])
        out.append({
            "sap": c["sap"], "name": c["name"], "tier": c["tier"], "ret": c["ret"], "status": c["status"],
            "codes": rec["codes"] if rec else [], "how": rec["how"] if rec else "",
            "n": rec["n"] if rec else 0, "ready": rec["ready"] if rec else 0,
            "fit": rec["fit"] if rec else "אין אנשי קשר",
            "people": rec["people"] if rec else [],
        })
    return out


# ---------- output ----------

def write_xlsx(mig, path):
    H = PatternFill("solid", fgColor="1F3864")
    HF = Font(name="Arial", bold=True, color="FFFFFF")
    B = Font(name="Arial", size=11)
    thin = Side(style="thin", color="D9D9D9")
    BD = Border(left=thin, right=thin, top=thin, bottom=thin)
    wb = openpyxl.Workbook()

    ws = wb.active
    ws.title = "לקוחות"
    ws.sheet_view.rightToLeft = True
    cols = ["מספר סאפ", "לקוח", "רמה", "ריטיינר", "סטטוס", "קודי חברה", "הוצלב לפי", "אנשי קשר", "מוכנים", "מוכנות"]
    ws.append(cols)
    for c in mig:
        ws.append([c["sap"], c["name"], c["tier"], c["ret"], c["status"], ", ".join(c["codes"]), c["how"], c["n"], c["ready"], c["fit"]])

    ws2 = wb.create_sheet("אנשי קשר")
    ws2.sheet_view.rightToLeft = True
    cols2 = ["מספר סאפ", "לקוח", "רמה", "שם", "תפקיד", "מס חברה", "מס סניף", "סניף", "מספר עובד", "נייד", "מייל", "ת.ז", "מוכן להעברה"]
    ws2.append(cols2)
    for c in mig:
        for p in c["people"]:
            ws2.append([c["sap"], c["name"], c["tier"], p["name"], p["role"], p["code"], p["branch_no"], p["branch"],
                        p["emp"], p["phone"], p["email"], p["tz"], "כן" if p["ok"] else "לא"])

    for sheet, widths in ((ws, [12, 36, 8, 10, 9, 14, 14, 10, 8, 14]), (ws2, [12, 30, 8, 24, 14, 9, 8, 26, 10, 14, 32, 12, 12])):
        for cell in sheet[1]:
            cell.fill, cell.font, cell.border = H, HF, BD
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font, cell.border = B, BD
                if isinstance(cell.value, str) and cell.value.isdigit():
                    cell.number_format = "@"
        for i, w in enumerate(widths, 1):
            sheet.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    wb.save(path)


def main():
    p = argparse.ArgumentParser(description="הצלבה מול הייצוא הטרי של משה")
    p.add_argument("--month", default="2026-07", help="חודש דוח ההכנסות (clients_status_<month>.json)")
    p.add_argument("--tag", default="2026-09", help="תג לקבצי הפלט")
    p.add_argument("--branches", default=str(DATA / "branches.csv"))
    p.add_argument("--employees", default=str(DATA / "employees.csv"))
    p.add_argument("--revenue", default=None)
    args = p.parse_args()

    rev, nets = load_config()
    revenue = Path(args.revenue) if args.revenue else sorted((ROOT / "data").glob("הכנסות*.xlsx"))[-1]
    clients = json.load(open(OUT / f"clients_status_{args.month}.json", encoding="utf-8"))

    branches = load_branches(args.branches)
    employees = load_employees(args.employees)
    by_hp, by_name, branch_names = companies_index(branches)
    hps, names, of_sap = read_families(revenue, rev, nets)
    linked = attach_people(link(clients, hps, names, of_sap, by_hp, by_name), employees, branch_names)
    mig = migration_set(clients, linked)

    print("סניפים: %d (%d חברות) · עובדים: %d שורות" % (len(branches), len({b['code'] for b in branches}), len(employees)))
    print("לקוחות: %d · מקושרים: %d (%.0f%%)" % (len(clients), len(linked), 100 * len(linked) / len(clients)))
    for how in ("ח.פ", "ח.פ + שם", "שם", "שם + שם"):
        n = sum(1 for v in linked.values() if v["how"] == how)
        if n:
            print("    לפי %s: %d" % (how, n))
    miss = [c for c in clients if c["sap"] not in linked]
    print("  לא מקושרים: %d · ריטיינר ₪%s" % (len(miss), format(sum(c["ret"] for c in miss), ",")))
    print("\nלקוחות המעבר (בינוני/קטן/מוב, לא 'לא פעיל'): %d" % len(mig))
    for fit in ("מוכן", "חסר פרט", "אין אנשי קשר"):
        print("  %s: %d" % (fit, sum(1 for c in mig if c["fit"] == fit)))
    print("  אנשי קשר: %d · מוכנים: %d" % (sum(c["n"] for c in mig), sum(c["ready"] for c in mig)))

    (OUT / f"sap_to_company_{args.tag}.json").write_text(
        json.dumps({k: {x: v[x] for x in ("codes", "scope", "how", "tier", "name", "ret", "status")} for k, v in linked.items()},
                   ensure_ascii=False, indent=1), encoding="utf-8")
    (OUT / f"migration_{args.tag}.json").write_text(json.dumps(mig, ensure_ascii=False, indent=1), encoding="utf-8")
    write_xlsx(mig, OUT / f"Fizikal_מעבר_{args.tag}.xlsx")
    print("\nנכתבו: out/sap_to_company_%s.json · out/migration_%s.json · out/Fizikal_מעבר_%s.xlsx" % (args.tag, args.tag, args.tag))


if __name__ == "__main__":
    main()
