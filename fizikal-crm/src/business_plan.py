#!/usr/bin/env python3
"""
Fizikal — מודל תוכנית עסקית: הכנסה חוזרת חודשית, ספטמבר 2026 עד דצמבר 2027, בשלושה תרחישים.

    python3 src/business_plan.py

החלטת Chen (24/09/2026): בסיס ₪1.0M (ריטיינר + מודולים, אוגוסט 2026 בלי הקפיצות החד-פעמיות),
יעד ₪1.3M בדצמבר 2027 (+30%). כל ההנחות בגיליון "הנחות" (כחול = אפשר לשנות), והכל מחושב בנוסחאות.

נתוני הבסיס נלקחו מ-data/revenue_latest.xlsx (דוח ההכנסות, ינואר-אוגוסט 2026), שורות ריטיינר + מודולים
של סאפים 2xxxxx/3xxxxx בלבד (בלי שורות סיכום והתאמות חשבונאיות).
פלט: out/Fizikal_תוכנית_עסקית_2027.xlsx
"""

from datetime import date

import openpyxl
from openpyxl.chart import LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

from revenue_pipeline import OUT

F = "Arial"
BLUE, GREEN = "0000FF", "008000"
HEAD_FILL = PatternFill("solid", fgColor="1F4E78")
SEC_FILL = PatternFill("solid", fgColor="D9E1F2")
KEY_FILL = PatternFill("solid", fgColor="FFFF00")
TOT_FILL = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
NIS = '₪#,##0;(₪#,##0);"-"'
PCT = '0.0%;(0.0%);"-"'
PCT2 = '0.00%;(0.00%);"-"'
MON = 'mmm-yy'

SCEN = ["שמרני", "בסיס", "אגרסיבי"]
SCOL = {"שמרני": "B", "בסיס": "C", "אגרסיבי": "D"}

# ---- נקודת פתיחה (ערך אחד) ----
START = [  # key, label, value, fmt, note
    ("base_ret", "ריטיינר חודשי — בסיס", 820_000, NIS,
     "אוגוסט 2026 בפועל: ₪822K, כולל קפיצה חד-פעמית של עיריית רעננה (₪18K+). מעוגל כך שריטיינר + מודולים = ₪1.0M (החלטת Chen)."),
    ("base_mod", "מודולים חודשי — בסיס", 180_000, NIS,
     "אוגוסט 2026 בפועל: ₪190K, כולל ~₪20K חיובים חריגים (טן דאנס, רמת פולג, מגדלי דוד, רעננה). יולי: ₪175K."),
    ("target", "יעד הכנסה חוזרת — דצמבר 2027", 1_300_000, NIS, "החלטת Chen: +30% על ₪1.0M."),
    ("clients", "לקוחות פעילים", 407, "#,##0", "client_status אוגוסט 2026: 399 פעיל + 8 חדש."),
]

# ---- הנחות לפי תרחיש ----
SCEN_INPUTS = [  # section/None, key, label, (cons, base, aggr), fmt, note, key-assumption?
    ("לקוחות חדשים", None, None, None, None, None, False),
    (None, "new_before", "ריטיינר מלקוחות חדשים לחודש — היום", (7_200, 7_200, 7_200), NIS,
     "בפועל פברואר-אוגוסט 2026: ₪43K מ-101 לקוחות חדשים ב-6 חודשים (ממוצע ₪430 ללקוח).", False),
    (None, "new_after", "ריטיינר מלקוחות חדשים לחודש — אחרי המיתוג", (9_000, 12_000, 18_000), NIS,
     "הנחה: מיתוג מחדש + מחירון באתר מגדילים את המכירות החדשות.", True),
    (None, "new_start", "המיתוג משפיע החל מ-", (date(2027, 1, 1),) * 3, MON, "", False),
    ("נטישה", None, None, None, None, None, False),
    (None, "churn_base", "נטישה חודשית — היום", (0.0075, 0.0075, 0.0075), PCT2,
     "בפועל: 43 לקוחות עזבו ב-6 חודשים, ₪44K = כ-0.75% מההכנסה בחודש.", False),
    (None, "churn_target", "נטישה חודשית — אחרי תוכנית שימור", (0.0070, 0.0050, 0.0040), PCT2,
     "הנחה: תוכנית Customer Success ללקוחות בסיכון.", True),
    (None, "churn_start", "תוכנית השימור משפיעה החל מ-", (date(2027, 1, 1),) * 3, MON, "", False),
    ("לקוחות קיימים", None, None, None, None, None, False),
    (None, "expansion", "הגדלות נטו חודשיות (הגדלות פחות הקטנות)", (0.0035, 0.0045, 0.0055), PCT2,
     "בפועל: הגדלות ₪82K פחות הקטנות ₪54K ב-6 חודשים = כ-0.47% בחודש.", False),
    ("תמחור חדש (בחידוש חוזה)", None, None, None, None, None, False),
    (None, "price_start", "המחירון החדש מתחיל ב-", (date(2027, 1, 1),) * 3, MON, "", False),
    (None, "price_months", "חודשים עד שכל הבסיס מחדש", (12, 12, 12), "0",
     "חוזים שנתיים: כל חודש מתחדשת 1/12 מהבסיס.", False),
    (None, "price_share", "חלק מהריטיינר שעובר למחירון החדש", (0.50, 0.70, 0.85), PCT,
     "השאר: VIP/גדול במחיר מותאם אישית, או לקוחות שמחירם כבר מעל המחירון.", True),
    (None, "price_uplift", "העלאה ממוצעת למי שעובר", (0.10, 0.15, 0.20), PCT,
     "לכייל מול סימולציית המחירון החדש לפי לקוח — לפני שמתחייבים למספר.", True),
    (None, "price_churn", "נטישה בגלל המחיר (מתוך מי שעובר)", (0.08, 0.05, 0.03), PCT,
     "ברוב חברות ה-SaaS, העלאה של 10%-20% מגדילה את הנטישה ב-3%-8% מהלקוחות שהמחיר שלהם עולה.", True),
    ("מכירת מודולים", None, None, None, None, None, False),
    (None, "mod_start", "מכירה יזומה מתחילה ב-", (date(2026, 10, 1),) * 3, MON, "", False),
    (None, "mod_clients", "לקוחות שמוסיפים מודול בחודש", (2, 4, 6), "0",
     "79 לקוחות בינוני/גדול בלי אף מודול (ריטיינר ₪134K).", True),
    (None, "mod_price", "ערך מודול ממוצע לחודש", (500, 500, 500), NIS,
     "היום: ₪175K מודולים על 163 לקוחות = ~₪1,070 ללקוח. מודול ראשון נמוך מזה.", False),
    ("פיצ'רים חדשים (₪50-100 לחודש)", None, None, None, None, None, False),
    (None, "addon_start", "השקה ב-", (date(2027, 3, 1), date(2027, 1, 1), date(2026, 11, 1)), MON, "", False),
    (None, "addon_adopt", "אימוץ בסוף הרמפה (מתוך הלקוחות הפעילים)", (0.15, 0.25, 0.40), PCT, "", True),
    (None, "addon_price", "מחיר ממוצע לחודש", (75, 75, 90), NIS, "החלטת Chen: ₪50-100 לפיצ'ר.", False),
    (None, "addon_ramp", "חודשים עד אימוץ מלא", (12, 12, 12), "0", "", False),
]


def font(color="000000", bold=False, size=10):
    return Font(name=F, color=color, bold=bold, size=size)


def build_inputs(wb):
    ws = wb.active; ws.title = "הנחות"; ws.sheet_view.rightToLeft = True
    ws["A1"] = "Fizikal — תוכנית עסקית 2027: הנחות"; ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = "כחול = הנחה שאפשר לשנות · צהוב = הנחות מפתח · שחור = נוסחה. כל המודל מתעדכן אוטומטית."
    ws["A2"].font = font("595959")
    ref = {}
    r = 4
    ws.cell(r, 1, "נקודת פתיחה (אוגוסט 2026)").font = font(bold=True)
    for c in range(1, 7): ws.cell(r, c).fill = SEC_FILL
    r += 1
    for key, label, val, fmt, note in START:
        ws.cell(r, 1, label).font = font()
        c = ws.cell(r, 2, val); c.font = font(BLUE); c.number_format = fmt
        ws.cell(r, 6, note).font = font("595959")
        ref[key] = f"'הנחות'!$B${r}"
        r += 1
    ws.cell(r, 1, "סה\"כ הכנסה חוזרת — בסיס").font = font(bold=True)
    c = ws.cell(r, 2, f"=B{r-4}+B{r-3}"); c.font = font(bold=True); c.number_format = NIS
    ref["base_total"] = f"'הנחות'!$B${r}"
    r += 2
    for i, h in enumerate(["הנחה", *SCEN, "", "מקור / הערה"], 1):
        c = ws.cell(r, i, h); c.font = font("FFFFFF", True); c.fill = HEAD_FILL
    r += 1
    for sec, key, label, vals, fmt, note, keyflag in SCEN_INPUTS:
        if sec:
            ws.cell(r, 1, sec).font = font(bold=True)
            for c in range(1, 7): ws.cell(r, c).fill = SEC_FILL
            r += 1; continue
        ws.cell(r, 1, label).font = font()
        for j, v in enumerate(vals):
            c = ws.cell(r, 2 + j, v); c.font = font(BLUE); c.number_format = fmt
            if keyflag: c.fill = KEY_FILL
        ws.cell(r, 6, note).font = font("595959")
        ref[key] = r
        r += 1
    ws.column_dimensions["A"].width = 44
    for col in "BCD": ws.column_dimensions[col].width = 13
    ws.column_dimensions["E"].width = 2
    ws.column_dimensions["F"].width = 95
    ws.freeze_panes = "B4"
    return ref


def build_model(wb, ref):
    ws = wb.create_sheet("מודל חודשי"); ws.sheet_view.rightToLeft = True
    ws["A1"] = "הכנסה חוזרת חודשית (ריטיינר + מודולים + פיצ'רים חדשים), ₪"; ws["A1"].font = font(bold=True, size=14)
    N = 16                                           # ספטמבר 2026 → דצמבר 2027
    first, last = 3, 3 + N - 1                       # C..R
    ws.cell(3, 1, "חודש").font = font(bold=True)
    ws.cell(3, 2, date(2026, 8, 1))
    ws.cell(3, 3, "=DATE(YEAR(B3),MONTH(B3)+1,1)")
    for c in range(4, last + 1):
        ws.cell(3, c, f"=DATE(YEAR({L(c-1)}3),MONTH({L(c-1)}3)+1,1)")
    for c in range(2, last + 1):
        x = ws.cell(3, c); x.number_format = MON; x.font = font("FFFFFF", True); x.fill = HEAD_FILL
        x.alignment = Alignment(horizontal="center")
    ws.cell(3, 1).font = font("FFFFFF", True); ws.cell(3, 1).fill = HEAD_FILL
    ws.cell(4, 2, "בסיס").font = font("595959"); ws.cell(4, 2).alignment = Alignment(horizontal="center")

    def P(key, s):                                   # הפניה להנחת תרחיש
        return f"'הנחות'!${SCOL[s]}${ref[key]}"

    blocks = {}
    r = 6
    for s in SCEN:
        rows = {}
        ws.cell(r, 1, f"תרחיש {s}").font = font("FFFFFF", True, 12)
        for c in range(1, last + 1): ws.cell(r, c).fill = HEAD_FILL
        r += 1
        labels = [
            ("ret_open", "ריטיינר — פתיחה"), ("new", "+ לקוחות חדשים"), ("churn", "− נטישה"),
            ("exp", "+ הגדלות נטו"), ("up", "+ העלאת מחיר בחידוש"), ("pchurn", "− נטישה בגלל מחיר"),
            ("ret_close", "ריטיינר — סגירה"), ("mod_open", "מודולים — פתיחה"), ("mod_new", "+ מכירת מודולים"),
            ("mod_churn", "− נטישת מודולים"), ("mod_close", "מודולים — סגירה"), ("addon", "פיצ'רים חדשים"),
            ("total", "סה\"כ הכנסה חוזרת"), ("path", "מסלול ליעד"), ("gap", "פער מהמסלול"), ("growth", "צמיחה מהבסיס"),
        ]
        for k, lab in labels:
            rows[k] = r; ws.cell(r, 1, lab).font = font(bold=k in ("ret_close", "mod_close", "total"))
            r += 1
        R = rows
        # עמודת בסיס (אוגוסט 2026)
        ws.cell(R["ret_close"], 2, f"={ref['base_ret']}").font = font(GREEN)
        ws.cell(R["mod_close"], 2, f"={ref['base_mod']}").font = font(GREEN)
        ws.cell(R["total"], 2, f"=B{R['ret_close']}+B{R['mod_close']}")
        ws.cell(R["path"], 2, f"={ref['base_total']}").font = font(GREEN)
        chunk = f"{ref['base_ret']}*{P('price_share', s)}/{P('price_months', s)}"
        for c in range(first, last + 1):
            X, Pv = L(c), L(c - 1)
            d = f"{X}$3"
            in_price = f"AND({d}>={P('price_start', s)},{d}<DATE(YEAR({P('price_start', s)}),MONTH({P('price_start', s)})+{P('price_months', s)},1))"
            churn_rate = f"IF({d}>={P('churn_start', s)},{P('churn_target', s)},{P('churn_base', s)})"
            months_live = f"((YEAR({d})-YEAR({P('addon_start', s)}))*12+MONTH({d})-MONTH({P('addon_start', s)})+1)"
            f = {
                "ret_open": f"={Pv}{R['ret_close']}",
                "new": f"=IF({d}>={P('new_start', s)},{P('new_after', s)},{P('new_before', s)})",
                "churn": f"=-{X}{R['ret_open']}*{churn_rate}",
                "exp": f"={X}{R['ret_open']}*{P('expansion', s)}",
                "up": f"=IF({in_price},{chunk}*(1-{P('price_churn', s)})*{P('price_uplift', s)},0)",
                "pchurn": f"=-IF({in_price},{chunk}*{P('price_churn', s)},0)",
                "ret_close": f"=SUM({X}{R['ret_open']}:{X}{R['pchurn']})",
                "mod_open": f"={Pv}{R['mod_close']}",
                "mod_new": f"=IF({d}>={P('mod_start', s)},{P('mod_clients', s)}*{P('mod_price', s)},0)",
                "mod_churn": f"=-{X}{R['mod_open']}*{churn_rate}",
                "mod_close": f"=SUM({X}{R['mod_open']}:{X}{R['mod_churn']})",
                "addon": f"=IF({d}>={P('addon_start', s)},{ref['clients']}*{P('addon_adopt', s)}*MIN(1,{months_live}/{P('addon_ramp', s)})*{P('addon_price', s)},0)",
                "total": f"={X}{R['ret_close']}+{X}{R['mod_close']}+{X}{R['addon']}",
                "path": f"={ref['base_total']}+({ref['target']}-{ref['base_total']})*{c - 2}/{N}",
                "gap": f"={X}{R['total']}-{X}{R['path']}",
                "growth": f"={X}{R['total']}/{ref['base_total']}-1",
            }
            for k, v in f.items():
                ws.cell(R[k], c, v)
        for k, rr in R.items():
            for c in range(2, last + 1):
                cell = ws.cell(rr, c); cell.number_format = PCT if k == "growth" else NIS
                if cell.font.color is None or cell.font.color.rgb in (None, "FF000000"):
                    cell.font = font(bold=k in ("ret_close", "mod_close", "total"))
                if k == "total":
                    cell.fill = TOT_FILL; cell.border = Border(top=THIN, bottom=THIN)
            if k == "total":
                ws.cell(rr, 1).fill = TOT_FILL
        blocks[s] = R
        r += 1
    ws.column_dimensions["A"].width = 26
    for c in range(2, last + 1): ws.column_dimensions[L(c)].width = 11.5
    ws.freeze_panes = "B4"
    return ws, blocks, first, last


def build_summary(wb, model, blocks, first, last, ref):
    ws = wb.create_sheet("סיכום", 0); ws.sheet_view.rightToLeft = True
    M = "'מודל חודשי'!"
    ws["A1"] = "Fizikal — תוכנית עסקית: מ-₪1.0M ל-₪1.3M הכנסה חוזרת בחודש עד דצמבר 2027"
    ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = "הכנסה חוזרת = ריטיינר + מודולים + פיצ'רים חדשים. לא כולל חומרה, פיתוח, SMS, התקנות ושותפויות. ההנחות בגיליון \"הנחות\"."
    ws["A2"].font = font("595959")
    Dc = L(last)
    hdr = ["", *SCEN]
    r = 4
    for i, h in enumerate(hdr, 1):
        c = ws.cell(r, i, h); c.font = font("FFFFFF", True); c.fill = HEAD_FILL; c.alignment = Alignment(horizontal="center")
    lines = [
        ("הכנסה חוזרת — אוגוסט 2026", lambda R: f"={M}B{R['total']}", NIS),
        ("הכנסה חוזרת — דצמבר 2027", lambda R: f"={M}{Dc}{R['total']}", NIS),
        ("צמיחה", lambda R: f"={M}{Dc}{R['growth']}", PCT),
        ("יעד", lambda R: f"={ref['target']}", NIS),
        ("פער ליעד (שלילי = חסר)", lambda R: f"={M}{Dc}{R['total']}-{ref['target']}", NIS),
        ("מגיע ליעד?", lambda R: f"=IF({M}{Dc}{R['total']}>={ref['target']},\"כן\",\"לא\")", "@"),
    ]
    for lab, fn, fmt in lines:
        r += 1
        ws.cell(r, 1, lab).font = font(bold=lab.startswith("הכנסה חוזרת — דצמבר"))
        for j, s in enumerate(SCEN):
            c = ws.cell(r, 2 + j, fn(blocks[s])); c.number_format = fmt; c.font = font(GREEN)
            c.alignment = Alignment(horizontal="center")
    r += 2
    ws.cell(r, 1, "מאיפה מגיע השינוי (סכום התנועות החודשיות, ספטמבר 2026 – דצמבר 2027)").font = font(bold=True)
    for c in range(1, 5): ws.cell(r, c).fill = SEC_FILL
    bridge = [("לקוחות חדשים", ["new"]), ("נטישה", ["churn", "mod_churn"]), ("הגדלות נטו", ["exp"]),
              ("העלאת מחיר בחידוש", ["up"]), ("נטישה בגלל מחיר", ["pchurn"]), ("מכירת מודולים", ["mod_new"])]
    b0 = r + 1
    for lab, keys in bridge:
        r += 1
        ws.cell(r, 1, lab).font = font()
        for j, s in enumerate(SCEN):
            R = blocks[s]
            f = "=" + "+".join(f"SUM({M}{L(first)}{R[k]}:{Dc}{R[k]})" for k in keys)
            c = ws.cell(r, 2 + j, f); c.number_format = NIS; c.font = font(GREEN)
    r += 1
    ws.cell(r, 1, "פיצ'רים חדשים (דצמבר 2027)").font = font()
    for j, s in enumerate(SCEN):
        c = ws.cell(r, 2 + j, f"={M}{Dc}{blocks[s]['addon']}"); c.number_format = NIS; c.font = font(GREEN)
    r += 1
    ws.cell(r, 1, "סה\"כ שינוי").font = font(bold=True)
    for j in range(3):
        col = L(2 + j)
        c = ws.cell(r, 2 + j, f"=SUM({col}{b0}:{col}{r-1})"); c.number_format = NIS; c.font = font(bold=True)
        c.fill = TOT_FILL
    ws.cell(r, 1).fill = TOT_FILL
    r += 1
    ws.cell(r, 1, "בדיקה: פתיחה + שינוי − דצמבר (צריך להיות 0)").font = font("595959")
    for j in range(3):
        col = L(2 + j)
        c = ws.cell(r, 2 + j, f"={col}5+{col}{r-1}-{col}6"); c.number_format = NIS; c.font = font("595959")
    for col, w in zip("ABCD", (52, 15, 15, 15)): ws.column_dimensions[col].width = w

    ch = LineChart(); ch.title = "הכנסה חוזרת חודשית מול מסלול היעד"; ch.height = 9; ch.width = 22
    ch.y_axis.title = "₪"; ch.y_axis.numFmt = '₪#,##0'; ch.x_axis.number_format = MON
    ch.y_axis.majorGridlines.spPr = None
    colors = {"שמרני": "EB6834", "בסיס": "2A78D6", "אגרסיבי": "1BAF7A"}
    for s in SCEN:
        R = blocks[s]
        ch.add_data(Reference(model, min_col=2, max_col=last, min_row=R["total"]), from_rows=True, titles_from_data=False)
        ser = ch.series[-1]; ser.tx = None
        from openpyxl.chart.series import SeriesLabel
        ser.tx = SeriesLabel(v=s); ser.graphicalProperties.line.solidFill = colors[s]; ser.graphicalProperties.line.width = 25400
        ser.smooth = False
    R = blocks["בסיס"]
    ch.add_data(Reference(model, min_col=2, max_col=last, min_row=R["path"]), from_rows=True, titles_from_data=False)
    from openpyxl.chart.series import SeriesLabel
    ser = ch.series[-1]; ser.tx = SeriesLabel(v="מסלול ליעד")
    ser.graphicalProperties.line.solidFill = "8C8C8C"; ser.graphicalProperties.line.dashStyle = "dash"; ser.graphicalProperties.line.width = 19050
    ch.set_categories(Reference(model, min_col=2, max_col=last, min_row=3))
    ch.legend.position = "b"
    ws.add_chart(ch, f"A{r + 3}")


def build_actuals(wb):
    ws = wb.create_sheet("בפועל 2026"); ws.sheet_view.rightToLeft = True
    ws["A1"] = "בפועל ינואר-אוגוסט 2026 (₪)"; ws["A1"].font = font(bold=True, size=14)
    ws["A2"] = ("מקור: דוח ההכנסות המעודכן (data/revenue_latest.xlsx), שורות 'הכנסות ריטיינר' + 'הכנסות מודולים' "
                "של סאפים 2xxxxx ו-3xxxxx. בלי שורות סיכום, התאמות חשבונאיות ותיקוני תקופות קודמות.")
    ws["A2"].font = font("595959")
    months = [date(2026, m, 1) for m in range(1, 9)]
    ret = [811_800, 801_600, 800_800, 832_800, 807_200, 810_100, 822_500, 822_500]
    mod = [168_800, 173_500, 173_600, 170_000, 173_700, 173_600, 175_300, 190_300]
    act = [399, 410, 425, 437, 460, 476, 468, 464]
    for i, h in enumerate(["חודש", "ריטיינר", "מודולים", "סה\"כ חוזר", "לקוחות משלמים"], 1):
        c = ws.cell(4, i, h); c.font = font("FFFFFF", True); c.fill = HEAD_FILL
    for k, m in enumerate(months):
        rr = 5 + k
        ws.cell(rr, 1, m).number_format = MON
        ws.cell(rr, 2, ret[k]).number_format = NIS
        ws.cell(rr, 3, mod[k]).number_format = NIS
        ws.cell(rr, 4, f"=B{rr}+C{rr}").number_format = NIS
        ws.cell(rr, 5, act[k]).number_format = "#,##0"
        for c in (2, 3, 5): ws.cell(rr, c).font = font(BLUE)
    ws.cell(12, 4).comment = Comment("כולל ~₪35K חיובים חריגים באוגוסט: עיריית רעננה, טן דאנס, רמת פולג, מגדלי דוד. לבדוק עם הנהלת החשבונות.", "Claude")
    ws.cell(14, 1, "שינוי ינואר-פברואר → יולי-אוגוסט").font = font(bold=True)
    ws.cell(14, 4, "=(D11+D12)/2-(D5+D6)/2").number_format = NIS
    ws.cell(15, 1, "באחוזים").font = font()
    ws.cell(15, 4, "=((D11+D12)/2)/((D5+D6)/2)-1").number_format = PCT

    ws.cell(17, 1, "גשר: ינואר-פברואר → יולי-אוגוסט (ממוצע לחודש)").font = font(bold=True)
    for c in range(1, 4): ws.cell(17, c).fill = SEC_FILL
    for i, h in enumerate(["תנועה", "לקוחות", "₪ לחודש"], 1):
        c = ws.cell(18, i, h); c.font = font("FFFFFF", True); c.fill = HEAD_FILL
    bridge = [("לקוחות חדשים", 101, 43_100), ("לקוחות שעזבו", 43, -44_100),
              ("הגדלות", None, 81_500), ("הקטנות", None, -53_800)]
    for k, (lab, n, v) in enumerate(bridge):
        rr = 19 + k
        ws.cell(rr, 1, lab)
        if n: ws.cell(rr, 2, n).font = font(BLUE)
        c = ws.cell(rr, 3, v); c.number_format = NIS; c.font = font(BLUE)
    ws.cell(23, 1, "נטו").font = font(bold=True)
    ws.cell(23, 3, "=SUM(C19:C22)").number_format = NIS
    ws.cell(25, 1, "לקוח חדש ממוצע: ₪430 בחודש. לקוח שעזב ממוצע: ₪1,025 בחודש. מביאים קטנים, מאבדים גדולים.").font = font("595959")
    for col, w in zip("ABCDE", (34, 13, 13, 13, 15)): ws.column_dimensions[col].width = w


def main():
    wb = openpyxl.Workbook()
    ref = build_inputs(wb)
    model, blocks, first, last = build_model(wb, ref)
    build_summary(wb, model, blocks, first, last, ref)
    build_actuals(wb)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if c.value is not None and (c.font is None or c.font.name != F):
                    c.font = Font(name=F, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                                  color=c.font.color if c.font else None)
    wb.calculation.fullCalcOnLoad = True
    out = OUT / "Fizikal_תוכנית_עסקית_2027.xlsx"
    wb.save(out)
    print(out)


if __name__ == "__main__":
    main()
