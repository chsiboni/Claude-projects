#!/usr/bin/env python3
"""
Fizikal — כיסוי תעודות זהות לקראת המעבר לגרסה החדשה.

הגרסה החדשה דורשת ת.ז כדי להתחבר, ולכן השאלה המעשית היא לא כמה אנשים חסרים
אלא האם לכל חברה יש לפחות משתמש בכיר אחד שיכול להיכנס — מייל תקין וגם ת.ז תקינה.

    python3 src/id_coverage.py --file 'data/תעודות_זהות.xlsx'

הקובץ הוא פלט marketing_contacts.py אחרי שהמתכנת הוסיף לו את ההצלבה למערכת
(חברה בפיזיקל, סניף בפיזיקל, שם עובד, מספר עובד, ת.ז).
"""

import argparse
import re
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
RANK = {"הרשאת על": 3, "ניהול חברה": 2, "מנהל": 1}
TIER_ORDER = {"VIP": 0, "גדול": 1, "בינוני": 2, "קטן": 3, "מוב": 4}


def digits(v):
    return re.sub(r"\D", "", str(v)) if v is not None else ""


def valid_id(s):
    """
    ספרת ביקורת של ת.ז ישראלית.

    בלי הבדיקה הזו נספרות כתקינות גם רשומות זבל מובהקות שיש בקובץ
    ('111111111', '987987987') וגם מספרים קטועים.
    """
    # '0' הוא הערך שהמערכת מחזירה כשלא נמצאה ת.ז, והוא דווקא *עובר* את ספרת
    # הביקורת (000000000), ולכן חייבים לפסול אותו במפורש לפני החישוב.
    if not s.isdigit() or len(s) > 9 or not s.strip("0"):
        return False
    s, total = s.zfill(9), 0
    for i, ch in enumerate(s):
        d = int(ch) * (1 if i % 2 == 0 else 2)
        total += d if d < 10 else d - 9
    return total % 10 == 0


def valid_email(e):
    return bool(re.fullmatch(r"[^@\s]+@[^@\s]+\.[^@\s]+", (e or "").strip().lower()))


def read(path):
    wb = openpyxl.load_workbook(path, data_only=True)

    # הת.ז יושבת בגיליון המנוכה כפילויות, אחת לאדם. ממופה לפי הנייד כדי להחזיר
    # אותה לגיליון המלא — אחרת חברה שכל אנשיה מנוהלים בידי אדם שנוכה מהגיליון
    # תיראה חשופה בטעות.
    ws = wb["דיוור"]
    idx = {str(c.value).strip(): n for n, c in enumerate(ws[2]) if c.value}
    id_by_phone = {}
    for r in ws.iter_rows(min_row=3, values_only=True):
        phone = digits(r[idx["נייד"]])
        if phone:
            id_by_phone[phone] = digits(r[idx["ת.ז"]])

    ws2 = wb["כל אנשי הקשר"]
    idx2 = {str(c.value).strip(): n for n, c in enumerate(ws2[2]) if c.value}
    people = []
    for r in ws2.iter_rows(min_row=3, values_only=True):
        if str(r[idx2["הוחרג מדיוור"]] or "").strip():
            continue                                   # פנימי ובדיקה
        people.append({
            "name": str(r[idx2["שם"]] or "").strip(),
            "company": str(r[idx2["שם חברה"]] or "").strip() or "(ללא שם חברה)",
            "client": str(r[idx2["לקוח"]] or "").strip(),
            "tier": str(r[idx2["Tier"]] or "").strip(),
            "role": str(r[idx2["תפקיד"]] or "").strip(),
            "email": valid_email(str(r[idx2["מייל"]] or "")),
            "id": valid_id(id_by_phone.get(digits(r[idx2["נייד"]]), "")),
        })
    return people


def by_company(people):
    comp = defaultdict(lambda: {"tier": "", "client": "", "n": 0, "ready": 0,
                                "mail_only": 0, "id_only": 0, "neither": 0,
                                "top": 0, "top_ready": 0, "who": None})
    for p in people:
        c = comp[p["company"]]
        c["n"] += 1
        c["tier"] = c["tier"] or p["tier"]
        c["client"] = c["client"] or p["client"]
        rank = RANK.get(p["role"], 0)
        c["top"] = max(c["top"], rank)
        if p["email"] and p["id"]:
            c["ready"] += 1
            if rank >= c["top_ready"]:
                c["top_ready"], c["who"] = rank, (p["name"], p["role"])
        elif p["email"]:
            c["mail_only"] += 1
        elif p["id"]:
            c["id_only"] += 1
        else:
            c["neither"] += 1
    return comp


def main():
    p = argparse.ArgumentParser(description="כיסוי ת.ז לקראת הגרסה החדשה")
    p.add_argument("--file", default="data/תעודות_זהות.xlsx")
    args = p.parse_args()

    people = read(ROOT / args.file if not Path(args.file).is_absolute() else args.file)
    comp = by_company(people)

    n = len(people)
    have_id = sum(1 for x in people if x["id"])
    ready = sum(1 for x in people if x["email"] and x["id"])
    print("אנשי קשר (בלי פנימי/בדיקה): %d" % n)
    print("  עם ת.ז תקינה:        %5d  (%.1f%%)" % (have_id, 100 * have_id / n))
    print("  בלי ת.ז שמישה:       %5d  (%.1f%%)" % (n - have_id, 100 * (n - have_id) / n))
    print("  מוכנים (מייל + ת.ז): %5d  (%.1f%%)" % (ready, 100 * ready / n))

    covered = {k: v for k, v in comp.items() if v["ready"]}
    exposed = {k: v for k, v in comp.items() if not v["ready"]}
    print("\nחברות: %d" % len(comp))
    print("  עם לפחות משתמש מוכן אחד: %3d  (%.0f%%)" % (len(covered), 100 * len(covered) / len(comp)))
    print("  חשופות לגמרי:            %3d  (%.0f%%)" % (len(exposed), 100 * len(exposed) / len(comp)))

    lvl = defaultdict(int)
    for v in covered.values():
        lvl[v["who"][1]] += 1
    print("\n  הדרג הבכיר ביותר שמוכן, בכל חברה מכוסה:")
    for role in sorted(lvl, key=lambda r: -RANK.get(r, 0)):
        print("    %-14s %3d" % (role, lvl[role]))
    gap = [v for v in covered.values() if v["top_ready"] < v["top"]]
    print("    מהן שדווקא הבכיר ביותר לא יכול להיכנס: %d" % len(gap))

    print("\n  החברות החשופות — ומה חסר בכל אחת:")
    print("  %-34s %-8s %5s %8s %7s %6s" % ("חברה", "Tier", "אנשים", "רק מייל", "רק ת.ז", "כלום"))
    for k, v in sorted(exposed.items(), key=lambda x: (TIER_ORDER.get(x[1]["tier"], 9), -x[1]["n"])):
        print("  %-34s %-8s %5d %8d %7d %6d"
              % (k[:33], v["tier"] or "—", v["n"], v["mail_only"], v["id_only"], v["neither"]))


if __name__ == "__main__":
    main()
