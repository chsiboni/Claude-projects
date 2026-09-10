#!/usr/bin/env python3
"""
Fizikal CRM — pipeline דוח ההכנסות.

קורא את דוח הנהלת החשבונות, מסנן שורות סיכום, מקבץ לרמת גוף משפטי (ח.פ)
תוך איחוד רשתות, מסווג Tier לפי הריטיינר החודשי, ומייצר קובץ לייבוא ל-monday.

    python3 src/revenue_pipeline.py --inspect                      # לאמת פריסת דוח חדש
    python3 src/revenue_pipeline.py --month 2026-06                # הרצה מלאה

כל ההחלטות העסקיות יושבות ב-config/ ולא בקוד. אם משהו כאן נראה שרירותי —
ההסבר נמצא בשדות ה-_comment / _note בקונפיג.
"""

import argparse
import json
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
CONFIG = ROOT / "config"
DATA = ROOT / "data"
OUT = ROOT / "out"

CATS = ["ret", "mod", "sms", "sms1", "train", "dev", "hw", "tags", "tech", "ship", "other"]


def load_config():
    rev = json.loads((CONFIG / "revenue.json").read_text(encoding="utf-8"))
    nets = json.loads((CONFIG / "nets.json").read_text(encoding="utf-8"))["nets"]
    return rev, nets


# ---------- normalisation helpers ----------

def s(v):
    return str(v).strip() if v is not None else ""


def digits(v):
    return re.sub(r"\D", "", str(v)) if v is not None else ""


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


# ---------- classification ----------

def category_of(detail, categories, ret_exclude):
    """
    סעיף ההכנסה → קטגוריה.

    'הכנסות ריטיינר תקופות קודמות' הוא חיוב או זיכוי רטרו ולכן נכנס ל-other, לא ל-ret:
    ה-Tier אמור לשקף את המצב הנוכחי, וסעיף כזה מכיל גם סכומים שליליים על העבר.
    """
    if detail in categories:
        return categories[detail]
    if "ריטיינר" in detail and not any(x in detail for x in ret_exclude):
        return "ret"
    return "other"


def family_key(name, hp, sap, nets):
    """
    מפתח הקיבוץ. רשתות גדולות לא חולקות ח.פ — לכל סניף ח.פ נפרד (ספייס = 14 ח.פ שונים),
    ואין בדוח שום עמודה שמזהה רשת. לכן איחוד רשתות נעשה לפי רשימה ידנית מאושרת.
    """
    for prefix, net in nets.items():
        if name.startswith(prefix):
            return ("net", net)
    return ("hp", hp) if hp else ("sap", sap)


def tier_of(retainer, tiers):
    """
    Tier נקבע לפי הריטיינר החודשי בלבד. הסדר קריטי: 'מוב' (בדיוק 182) נבדק אחרי
    הטווחים הגבוהים אבל לפני 'קטן', אחרת 182 ייפול ל'קטן'.
    """
    for t in tiers:
        if "exact" in t:
            if abs(retainer - t["exact"]) < 1:
                return t["name"]
            continue
        lo, hi = t.get("min"), t.get("max")
        if lo is not None and retainer < lo:
            continue
        if hi is not None and retainer >= hi:
            continue
        return t["name"]
    return None  # ריטיינר 0 — לא לקוח משלם החודש


# ---------- reading ----------

def canonical_name(names):
    """
    שם מייצג למשפחה. 16 מ-352 משפחות ח.פ מופיעות תחת יותר משם אחד — איות שונה
    בשורות שונות, לפעמים עם סיומת כמו 'סיים פעילות'. בוחרים את הנפוץ, ובתיקו את הקצר,
    כדי שהבחירה לא תהיה תלויה בסדר השורות בדוח.
    """
    if not names:
        return ""
    return min(names, key=lambda n: (-names[n], len(n), n))


def disambiguate(clients):
    """
    שני גופים משפטיים יכולים לחלוק שם מסחרי (למשל 'מנהל קהילתי יובלים (ער)' תחת שני ח.פ).
    שורות כאלה בבורד נראות ככפילות, ולכן מצמידים להן את הח.פ כדי שיהיו נפרדות בעין.
    """
    seen = Counter(c["name"] for c in clients)
    for c in clients:
        if seen[c["name"]] > 1:
            tail = f"ח.פ {c['hp']}" if c["hp"] else f"סאפ {c['_sap']}"
            c["name"] = f"{c['name']} · {tail}"
    return clients


def resolve_hp_by_sap(rows, layout):
    """
    ח.פ לכל סאפ, מכל השורות שלו. סאפ→ח.פ הוא 1:1 ולכן זה בטוח.

    בלי זה, שורה שבה עמודת ח.פ ריקה נופלת לקיבוץ לפי סאפ ומייצרת 'לקוח פנטום'
    שהוא בעצם שבר של לקוח קיים.

    לדוח שתי עמודות ח.פ, והן לא שוות בערכן: העמודה הראשית קובעת, והמשנית משמשת
    רק כשאין לסאפ שום ערך בראשית. הסדר הזה חשוב — המשנית סותרת את הראשית בכמה
    סאפים בודדים (נראה כמו שגיאות הקלדה), אבל היא מוסיפה ח.פ למעל מאה לקוחות
    שאחרת היו מקובצים לפי סאפ ולא ניתנים להצלבה מול הסניפים.
    """
    primary, secondary, conflicts = {}, {}, {}
    for row in rows:
        sap = s(row[layout["col_sap"]])
        if not sap:
            continue
        hp = digits(row[layout["col_hp"]])
        # '0' הוא ערך זבל שמוקלד כשאין ח.פ — לא לתת לו לחסום את העמודה המשנית
        if hp.strip("0"):
            if sap in primary and primary[sap] != hp:
                conflicts.setdefault(sap, {primary[sap]}).add(hp)
            primary.setdefault(sap, hp)
        alt_col = layout.get("col_hp_alt")
        if alt_col is not None:
            alt = digits(row[alt_col])
            if alt.strip("0"):
                secondary.setdefault(sap, alt)

    hp_by_sap = dict(secondary)
    hp_by_sap.update(primary)   # הראשית דורסת את המשנית
    recovered = {sap for sap in secondary if sap not in primary}
    return hp_by_sap, conflicts, recovered


def open_sheet(path):
    if not path.exists():
        sys.exit(
            f"לא נמצא: {path}\n"
            f"הקבצים הרגישים לא בגיט — צריך להעלות את הדוח ל-{DATA}/ בכל סשן."
        )
    return openpyxl.load_workbook(path, data_only=True).active


def inspect(path, layout):
    """מדפיס את שורת הכותרות ושורות דוגמה, כדי לאמת שהפריסה בקונפיג עדיין נכונה."""
    ws = open_sheet(path)
    rows = list(ws.iter_rows(min_row=1, max_row=layout["first_data_row"] + 3, values_only=True))
    header = rows[layout["header_row"] - 1]
    print(f"גיליון: {ws.title} · {ws.max_row} שורות · {ws.max_column} עמודות\n")
    print(f"--- שורת כותרות ({layout['header_row']}) ---")
    for i, cell in enumerate(header):
        print(f"  [{i:>2}] {s(cell) or '(ריק)'}")
    print(f"\n--- 3 שורות נתונים ראשונות (מ-{layout['first_data_row']}) ---")
    for r in rows[layout["first_data_row"] - 1:]:
        print("  " + " | ".join(f"[{i}]{s(c)}" for i, c in enumerate(r) if s(c))[:400])
    print("\nלאמת מול config/revenue.json → layout. אם ההתאמה שבורה — לעדכן שם, לא בקוד.")


def build_clients(path, month, rev, nets):
    layout = rev["layout"]
    if month not in layout["month_columns"]:
        sys.exit(
            f"החודש {month} לא מוגדר בקונפיג. חודשים זמינים: "
            f"{', '.join(layout['month_columns'])}\n"
            f"דוח חדש? להריץ --inspect ולהוסיף את העמודה ל-month_columns."
        )
    mcol = layout["month_columns"][month]
    sap_re = re.compile(rev["client_sap_pattern"])
    skip_names = set(rev["skip_names"])
    skip_saps = set(rev["skip_saps"])
    skip_details = set(rev["skip_details"])

    ws = open_sheet(path)
    rows = [r for r in ws.iter_rows(min_row=layout["first_data_row"], values_only=True) if any(r)]
    hp_by_sap, hp_conflicts, hp_from_alt = resolve_hp_by_sap(rows, layout)

    fam = defaultdict(lambda: {"name": "", "hp": "", "saps": set(),
                               "names": Counter(), "m": dict.fromkeys(CATS, 0.0)})
    stats = defaultdict(int)

    for row in rows:
        stats["rows_read"] += 1
        sap = s(row[layout["col_sap"]])
        name = s(row[layout["col_name"]])
        detail = s(row[layout["col_detail"]])

        # פילטר הסאפ הוא הקו הראשון — הוא זורק גם את הטבלה המסכמת בתחתית הגיליון
        if not sap_re.fullmatch(sap) or sap in skip_saps:
            stats["skip_not_client_sap"] += 1
            continue
        if not name:
            stats["skip_no_name"] += 1
            continue
        if name in skip_names or any(name.startswith(p) for p in rev["skip_name_prefixes"]):
            stats["skip_summary_name"] += 1
            continue
        if detail in skip_details:
            stats["skip_summary_detail"] += 1
            continue

        stats["rows_kept"] += 1
        hp = hp_by_sap.get(sap, "")
        if not digits(row[layout["col_hp"]]) and hp:
            stats["hp_recovered"] += 1
        kind, key = family_key(name, hp, sap, nets)
        f = fam[(kind, key)]
        f["saps"].add(sap)
        if hp and not f["hp"]:
            f["hp"] = hp
        if kind == "net":
            f["name"] = key
        else:
            f["names"][name] += 1
        f["m"][category_of(detail, rev["categories"], rev["retainer_categories_exclude"])] += num(row[mcol])

    for f in fam.values():
        if not f["name"]:
            f["name"] = canonical_name(f["names"])

    clients = []
    for f in fam.values():
        tier = tier_of(f["m"]["ret"], rev["tiers"])
        if not tier:
            stats["dropped_zero_retainer"] += 1
            continue
        m = f["m"]
        clients.append({
            "tier": tier,
            "name": f["name"],
            "hp": f["hp"],
            "nsap": len(f["saps"]),
            "_sap": min(f["saps"]),
            "ret": round(m["ret"]),
            "mod": round(m["mod"]),
            "sms": round(m["sms"] + m["sms1"]),
            "train": round(m["train"]),
            "dev": round(m["dev"]),
            "hw": round(m["hw"]),
            "tags": round(m["tags"]),
            "other": round(m["tech"] + m["ship"] + m["other"]),
            "total": round(sum(m.values())),
        })

    disambiguate(clients)
    clients.sort(key=lambda c: -c["ret"])
    stats["hp_conflicts"] = len(hp_conflicts)
    stats["hp_from_alt"] = len(hp_from_alt)
    return clients, stats, hp_conflicts


# ---------- writing ----------

COLUMNS = [
    ("Name", "name"), ("Tier", "tier"), ("ח.פ", "hp"), ("ריטיינר", "ret"),
    ("סה\"כ", "total"), ("מודולים", "mod"), ("SMS", "sms"),
    ("הדרכות והתקנה", "train"), ("פיתוח", "dev"), ("חומרה", "hw"),
    ("תגי קרבה", "tags"), ("אחר", "other"), ("מספר כרטיסי חיוב", "nsap"),
]


def write_outputs(clients, month, rev):
    OUT.mkdir(exist_ok=True)
    xlsx = OUT / f"clients_{month}.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "לקוחות"
    ws.sheet_view.rightToLeft = True
    # העמודה הראשונה חייבת להיקרא Name, אחרת monday מפילה כל שורה ב-'Name must not be empty'
    ws.append([title for title, _ in COLUMNS])
    for c in clients:
        ws.append([c[key] for _, key in COLUMNS])
    wb.save(xlsx)

    js = OUT / f"clients_{month}.json"
    js.write_text(json.dumps(clients, ensure_ascii=False, indent=1), encoding="utf-8")
    return xlsx, js


def report(clients, stats, month, rev):
    dist = defaultdict(int)
    for c in clients:
        dist[c["tier"]] += 1
    order = [t["name"] for t in rev["tiers"]]

    print(f"\n=== {month} ===")
    print(f"שורות בגיליון: {stats['rows_read']:>7,}")
    for k, label in [
        ("skip_not_client_sap", "  סוננו — לא סאפ לקוח"),
        ("skip_no_name", "  סוננו — בלי שם חברה"),
        ("skip_summary_name", "  סוננו — שורת סיכום"),
        ("skip_summary_detail", "  סוננו — פירוט מסכם"),
    ]:
        print(f"{label}: {stats[k]:>7,}")
    print(f"שורות לקוח נותרו: {stats['rows_kept']:>5,}")
    print(f"ח.פ שוחזר משורה אחרת: {stats['hp_recovered']:>3,}")
    print(f"ח.פ מהעמודה המשנית: {stats['hp_from_alt']:>5,} סאפים")
    print(f"נזרקו (ריטיינר 0): {stats['dropped_zero_retainer']:>5,}")

    print(f"\nלקוחות: {len(clients)}")
    for t in order:
        print(f"  {t:<8} {dist[t]:>4}")

    print(f"\nריטיינר חודשי:  ₪{sum(c['ret'] for c in clients):>12,}")
    print(f"סה\"כ החודש:     ₪{sum(c['total'] for c in clients):>12,}")

    expected = rev.get("expected_distribution_2026_06")
    if expected and month == "2026-06":
        actual = {**dist, "total": len(clients)}
        diffs = [f"{k}: {expected[k]}→{actual.get(k, 0)}" for k in expected if expected[k] != actual.get(k)]
        print("\n" + ("✅ תואם את ההתפלגות המאומתת של יוני 2026"
                      if not diffs else "⚠️  סטייה מההתפלגות המאומתת: " + " · ".join(diffs)))

    top = clients[:5]
    print("\n5 הגדולים לפי ריטיינר:")
    for c in top:
        print(f"  {c['tier']:<6} ₪{c['ret']:>9,}  {c['name']}  ({c['nsap']} כרטיסים)")


def main():
    p = argparse.ArgumentParser(description="pipeline דוח ההכנסות של Fizikal")
    p.add_argument("--file", default=None, help="נתיב לדוח (ברירת מחדל: data/הכנסות <חודש>.xlsx)")
    p.add_argument("--month", default=None, help='חודש הדוח, למשל 2026-06')
    p.add_argument("--inspect", action="store_true", help="להדפיס את פריסת הגיליון ולצאת")
    args = p.parse_args()

    rev, nets = load_config()

    if args.file:
        path = Path(args.file)
    else:
        candidates = sorted(DATA.glob("הכנסות*.xlsx"))
        if not candidates:
            sys.exit(f"לא נמצא דוח הכנסות ב-{DATA}/ — להעלות אותו לשם, או להעביר --file")
        path = candidates[-1]
        print(f"דוח: {path.name}")

    if args.inspect:
        inspect(path, rev["layout"])
        return

    if not args.month:
        sys.exit("חסר --month (למשל --month 2026-06). --inspect מראה אילו עמודות חודש קיימות.")

    clients, stats, hp_conflicts = build_clients(path, args.month, rev, nets)
    if hp_conflicts:
        print("⚠️  סאפ עם יותר מח.פ אחד — ההנחה של 1:1 נשברה, לבדוק:")
        for sap, hps in list(hp_conflicts.items())[:10]:
            print(f"     סאפ {sap}: {', '.join(sorted(hps))}")
    report(clients, stats, args.month, rev)
    xlsx, js = write_outputs(clients, args.month, rev)
    print(f"\nנכתב:\n  {xlsx.relative_to(ROOT)}  ← לייבוא ל-monday\n  {js.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
