"""רשימת שיחות לאימות פרטי לקוחות — שורה ללקוח, 4 מועמדים ליצירת קשר לפי בכירות, וגיליון טופס מילוי ריק.
קלט: dump JSON.gz של ארבעה בורדים + collections.json / collections_missing.json מבורד הגבייה.
הרצה: python3 src/call_list_2026.py <dump.json.gz>
"""
import gzip,json,collections as C,re,sys
import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.worksheet.datavalidation import DataValidation
D=json.load(gzip.open(sys.argv[1] if len(sys.argv)>1 else '/tmp/claude-0/-home-user-Claude-projects/9bedc6a6-f60e-5cce-a499-d0c73e5871ea/scratchpad/dl/all.json.gz','rt',encoding='utf-8'))
COLM=json.load(open('collections_missing.json',encoding='utf-8')); COL=json.load(open('collections.json',encoding='utf-8'))
comp,br,onb,con=D['companies'],D['branches'],D['onboarding'],D['contacts']
norm=lambda s:(s or '').strip()
clean=lambda p: re.sub(r'\D','',p or '')
def okphone(p): return len(p) in(9,10) and p.startswith('0') and not p.startswith('00')
mob=lambda p: re.fullmatch(r'05\d{8}',p) is not None
BRC=C.defaultdict(list); SAPB=C.defaultdict(list)
for b in br:
    b['_cc']=norm(b['text_mm6v5fq7']); b['_bn']=norm(b['text_mm6vxmyx'])
    if b['_cc']: BRC[b['_cc']].append(b)
    if norm(b['text_mm6vwm4f']): SAPB[norm(b['text_mm6vwm4f'])].append(b)
CON=C.defaultdict(list)
for c in con: CON[norm(c['text_mm6v6myv'])].append(c)
COLSAP={x['sap']:x for x in COL}
RANK={'ניהול חברה':0,'הרשאת על':1}
def mg_rank(c):
    r=norm(c['dropdown_mm6v5453']); g=norm(c['dropdown_mm6v5svk'])
    if r in RANK: return RANK[r]
    if r=='מנהל' and g in('הנהלה','הכל','מנהל'): return 2
    if r=='מנהל': return 3
    return None
ROLE={0:'ניהול חברה',1:'הרשאת על',2:'מנהל (הנהלה)',3:'מנהל'}
def candidates(ccs,sap=''):
    out=[]; seen=set()
    ppl=[]
    for cc in ccs:
        for c in CON.get(cc,[]):
            rk=mg_rank(c); p=clean(c['text_mm6vqqpn'])
            if rk is None or not okphone(p): continue
            ppl.append((rk,0 if mob(p) else 1,norm(c['name']),ROLE[rk],p))
    for rk,m,n,r,p in sorted(ppl):
        if p in seen: continue
        seen.add(p); out.append((n,r,p))
        if len(out)==4: return out
    for cc in ccs:
        for b in BRC.get(cc,[]):
            p=clean(b['text_mm6vnk3m'])
            if okphone(p) and p not in seen:
                seen.add(p); out.append((norm(b['name']),'טלפון הסניף',p))
                if len(out)==4: return out
    if sap in COLSAP:
        p=clean(COLSAP[sap]['phone'])
        if okphone(p) and p not in seen: out.append((COLSAP[sap]['email'] or 'הנה"ח של הלקוח','איש הכספים (מבורד גבייה)',p))
    return out[:4]
tier_of=lambda g:{'VIP':'VIP','גדולים':'גדול','בינוניים':'בינוני','קטנים':'קטן','מוב':'מוב'}.get(g.replace(' - לא פעילים','').strip(),g)
rows=C.defaultdict(list); linked=set(); st=C.Counter()
def add(sheet,sap,name,state,hp,cands):
    r=[sap,name,state,hp]
    for i in range(4):
        r+= list(cands[i]) if i<len(cands) else ['','','']
    rows[sheet].append(r); st['cands_'+str(len(cands))]+=1
for co in comp:
    sap=norm(co['text_mm6ccg01']); state='לא פעיל 3+ חודשים' if co['color_mm6cyme4']=='לא פעיל' else co['color_mm6cyme4']
    ccs=list(dict.fromkeys([norm(s.get('text_mm76nbjb')) for s in co.get('subitems',[]) if norm(s.get('text_mm76nbjb'))] or [b['_cc'] for b in SAPB.get(sap,[])]))
    linked.update(ccs)
    add(tier_of(co['group']),sap,norm(co['name']),state,norm(co['text_mm5ebgdf']),candidates(ccs,sap))
LIVE={'אונבורדינג הסתיים','עלה לאוויר- ליווי חודש 1','ליווי חודש 2','באוויר - ממתין לסליקה','הדרכה 1','הדרכה 2','הדרכה עיכובים','ממתין להדרכה','שלבי הסבה (אופציונלי)','ממתין לסליקה (מכירות)','סיכום מכירה'}
seen_cc=set()
for o in onb:
    if not (o['date_mm2fvzsy']>='2026-05-01' or (not o['date_mm2fvzsy'] and o['status'] in LIVE and o['date4']>='2026-05-01')): continue
    if o['status']=='ביטל עסקה': continue
    cc=norm(o['text_mm2bb1bq'])
    if cc and cc in linked: continue
    if cc and cc in BRC: seen_cc.add(cc); add('לא מסווג',cc,norm(o['name']),'אונבורדינג מאי+','',candidates([cc]))
    else: add('לא מסווג',cc,norm(o['name']),'אונבורדינג מאי+'+('' if cc else ', בלי מספר חברה'),'',[])
act=C.defaultdict(list)
for b in br:
    if b['color_mm6vn3xc']=='פעיל' and b['_cc'] and b['_cc'] not in linked and b['_cc'] not in seen_cc: act[b['_cc']].append(b)
for cc,bs in act.items(): add('לא מסווג',cc,norm(bs[0]['text_mm6vcn1r']) or norm(bs[0]['name']),'פעיל במערכת, לא בדוח הכנסות',norm(bs[0]['text_mm6vx83h']),candidates([cc]))
CH=['פרופיט','פיטפורלס','גרייט שייפ','סוויפט','מטי מועדוני','מגדלי הים התיכון','בארבור','אייקון פיטנס']
for m in COLM:
    chain=next((c for c in CH if c in m['name']),''); p=clean(m['phone'])
    add('לא מסווג',m['sap'],norm(m['name']),'בורד גביה, לא בדוח אוגוסט'+(f' (כנראה סניף של רשת {chain})' if chain else ''),norm(m.get('cc','')),[(m['email'] or 'הנה"ח','איש הכספים (מבורד גבייה)',p)] if okphone(p) else [])
print(dict(st))
for k in ['VIP','גדול','בינוני','קטן','מוב','לא מסווג']: print(k,len(rows[k]),'| 4 cands',sum(1 for r in rows[k] if r[13]),'| 0 cands',sum(1 for r in rows[k] if not r[4]))
# ---------- write
F='Arial'; wb=openpyxl.Workbook(); wb.remove(wb.active)
thin=Side(style='thin',color='C9D1E0'); B=Border(top=thin,bottom=thin,left=thin,right=thin)
HEAD=['מספר לקוח','שם הלקוח','מצב אצלנו','ח.פ']+[f'{t} {i}' for i in (1,2,3,4) for t in ('שם','תפקיד','טלפון')]+['סטטוס שיחה','דיברנו עם (שם ותפקיד)','תאריך','הערות']
W=[11,30,20,12]+[18,16,13]*4+[14,24,11,30]
for k in ['VIP','גדול','בינוני','קטן','מוב','לא מסווג']:
    ws=wb.create_sheet(k); ws.sheet_view.rightToLeft=True
    for i,(h,w) in enumerate(zip(HEAD,W),1):
        c=ws.cell(row=1,column=i,value=h); c.font=Font(name=F,bold=True,color='FFFFFF')
        c.fill=PatternFill('solid',fgColor='1F2A44' if i<=4 else ('2E5A88' if i<=16 else '6B7280')); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=B
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    ws.row_dimensions[1].height=30
    for r,row in enumerate(rows[k],2):
        for i,v in enumerate(row,1):
            c=ws.cell(row=r,column=i,value=v or None); c.font=Font(name=F,size=10); c.border=B
            if i in(1,4,7,10,13,16): c.number_format='@'
        for i in range(len(row)+1,len(HEAD)+1): ws.cell(row=r,column=i).border=B
    n=len(rows[k])+1
    dv=DataValidation(type='list',formula1='"הושלם,לא הושג,סירב,העסק נסגר"',allow_blank=True); ws.add_data_validation(dv); dv.add(f'Q2:Q{n}')
    ws.freeze_panes='C2'; ws.auto_filter.ref=f'A1:{openpyxl.utils.get_column_letter(len(HEAD))}{max(n,2)}'
# fill form sheet
ws=wb.create_sheet('טופס מילוי'); ws.sheet_view.rightToLeft=True
FIELDS=['שם חברה','ח.פ','שם לחשבונית','מספר חברה','מספר סניף','שם סניף','סטטוס לקוח','שם מלא','תפקיד','ת.ז','מייל','טלפון','אומת בתאריך']
ws['A1']='טופס מילוי — העתק לכל לקוח (שם הקובץ = מספר הלקוח). שורה לכל איש הנהלה בכל סניף.'; ws['A1'].font=Font(name=F,bold=True,size=12)
ws['A2']='מספר לקוח:'; ws['B2']=''; ws['A2'].font=Font(name=F,bold=True)
for i,(h,w) in enumerate(zip(FIELDS,[26,12,26,10,9,22,12,24,14,12,28,13,13]),1):
    c=ws.cell(row=4,column=i,value=h); c.font=Font(name=F,bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F2A44'); c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True); c.border=B
    ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width=w
    for r in range(5,25): ws.cell(row=r,column=i).border=B
dv=DataValidation(type='list',formula1='"בעלים,מנהל סניף,הנהלה אחרת,קבלה,מנהל משמרת,אחר"',allow_blank=True); ws.add_data_validation(dv); dv.add('I5:I200')
dv2=DataValidation(type='list',formula1='"פעיל,לא פעיל"',allow_blank=True); ws.add_data_validation(dv2); dv2.add('G5:G200')
ws=wb.create_sheet('הסבר',0); ws.sheet_view.rightToLeft=True; ws.column_dimensions['A'].width=115
L=["רשימת שיחות לאימות פרטי לקוחות — Fizikal, 12/10/2026",
"גיליון לכל רמת לקוח לפי דוח ההכנסות אוגוסט 2026 (VIP · גדול · בינוני · קטן · מוב), ו'לא מסווג' למי שלא בדוח: אונבורדינג שעלה לאוויר ממאי, חברות פעילות במערכת בלי תשלום, וכרטיסים מבורד הגבייה.",
"שורה = לקוח. לכל לקוח עד 4 מועמדים ליצירת קשר, מהבכיר לפחות בכיר: ניהול חברה → הרשאת על → מנהל בהנהלה → מנהל → טלפון הסניף → איש הכספים. מתחילים ממועמד 1.",
"התהליך: מתקשרים, מבקשים את מי שמנהל את הפעילות (המנהל הבכיר ביותר שניתן להשיג). מולו ממלאים את 'טופס מילוי' (העתק לכל לקוח, שם הקובץ = מספר הלקוח): שורה לכל איש הנהלה בכל סניף. בסוף מעדכנים בשורת הלקוח: סטטוס שיחה, עם מי דיברנו, תאריך.",
"הנתונים שלנו (שמות, טלפונים) הם רק כדי להגיע לאדם הנכון. לא מעתיקים אותם לטופס. המנהל ממלא הכול מחדש.",
"סטטוס שיחה: הושלם / לא הושג (אחרי 3 ניסיונות בשעות שונות) / סירב / העסק נסגר."]
for i,t in enumerate(L,1):
    c=ws.cell(row=i,column=1,value=t); c.font=Font(name=F,bold=(i==1),size=13 if i==1 else 11); c.alignment=Alignment(wrap_text=True,vertical='top',horizontal='right'); ws.row_dimensions[i].height=22 if i==1 else 38
out='/home/user/Claude-projects/fizikal-crm/out/Fizikal_רשימת_שיחות_לאימות_2026-10.xlsx'; wb.save(out); print('saved')
