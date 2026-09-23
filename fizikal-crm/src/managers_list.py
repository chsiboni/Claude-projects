#!/usr/bin/env python3
"""
Fizikal — רשימת דיוור: איש קשר מנהלי אחד לכל חברה פעילה.

    python3 src/managers_list.py        # קלט: out/mgr_raw.json (ייצוא מ-monday, ראה למטה)

בקשת Chen (23/09/2026) עבור מנהלת השיווק: גיליון אחד, שורה לחברה, עמודות
שם · טלפון · מייל · רמת לקוח · חברה · סניף · דרג. כולל חברות שלא שויכו (דרג "לא משויך",
טלפון/מייל מבורד הגביה אם יש). "לא פעיל" (3+ חודשים בלי תשלום) — לא נכלל.

מי נבחר בחברה: ניהול חברה → מנהל → מנהל משמרת; בתיקו — מי שיש לו נייד+מייל, אחר כך הוותיק.

out/mgr_raw.json נוצר ב-execute_code של monday MCP: companies = בורד חברות 5102514504,
contacts = בורד אנשי קשר 5103469435 (רק שורות עם מספר סאפ), gviya = בורד גביה 5102318900.
"""
import json, re, collections as C
import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from revenue_pipeline import OUT

RANK = {'ניהול חברה': 3, 'מנהל': 2, 'מנהל משמרת': 1}
TIER = {'VIP': 0, 'גדול': 1, 'בינוני': 2, 'קטן': 3, 'מוב': 4}


def ph(p):
    p = re.sub(r'\D', '', p or '')
    if p.startswith('972'): p = '0' + p[3:]
    if len(p) == 9 and p[0] == '5': p = '0' + p
    return p if re.fullmatch(r'05\d{8}', p) else ''


def em(e):
    e = (e or '').strip().lower()
    return e if re.fullmatch(r'[^@\s,;]+@[^@\s,;]+\.[a-z]{2,}', e) else ''


def main():
    d = json.load(open(OUT / 'mgr_raw.json', encoding='utf-8'))
    by = C.defaultdict(list)
    for c in d['contacts']:
        for s in [x.strip() for x in c['text_mm6vke74'].split(',') if x.strip()]:
            by[s].append(c)
    gv = {g['text_mm69gg7g']: g for g in d['gviya'] if g['text_mm69gg7g']}

    def key(c):
        p, e = ph(c['text_mm6vqqpn']), em(c['text_mm6vwyzq'])
        return (-RANK.get(c['dropdown_mm6v5453'], 0), -(bool(p) and bool(e)), -bool(p), -bool(e),
                c['date_mm6v5er7'] or '9999', c['n'])

    rows = []
    for co in d['companies']:
        if co['color_mm6cyme4'] == 'לא פעיל':
            continue
        sap, tier = co['text_mm6ccg01'], co['color_mm5ek8b7']
        name = re.sub(r'\s*·\s*סאפ \d+$', '', co['n'])
        if by.get(sap):
            c = sorted(by[sap], key=key)[0]
            rows.append([c['n'], ph(c['text_mm6vqqpn']), em(c['text_mm6vwyzq']), tier, name, c['text_mm6vw0jh'], c['dropdown_mm6v5453']])
        else:
            g = gv.get(sap, {})
            rows.append(['', ph(g.get('phone_mm699f4', '')), em(g.get('email_mm6k2wy6', '')), tier, name, '', 'לא משויך'])
    rows.sort(key=lambda r: (TIER.get(r[3], 9), r[4]))

    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'מנהלים'; ws.sheet_view.rightToLeft = True
    ws.append(['שם', 'מספר טלפון', 'מייל', 'רמת לקוח', 'חברה', 'סניף', 'דרג'])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F4E78'); c.alignment = Alignment(horizontal='center')
    for r in rows:
        ws.append(r)
    for (c,) in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        c.number_format = '@'
    for i, w in enumerate([24, 14, 32, 10, 36, 30, 14], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    out = OUT / 'Fizikal_מנהלים_לדיוור_2026-09.xlsx'
    wb.save(out)
    print(out.name, len(rows), C.Counter(r[6] for r in rows))


if __name__ == '__main__':
    main()
