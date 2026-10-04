"""תוכנית עבודה 2027 — אקסל היעדים לפי רבעון, שתי טבלאות: צמיחה 15% וצמיחה 30%.

העתק של המודל מדף ה-HTML (plan2027) בפייתון, ובניית חוברת עבודה מבוססת נוסחאות:
  הנחות        — כל הקלט (כחול על צהוב): קבועים, מחירון, ושני בלוקי תרחיש
  מודל 15/30   — 16 חודשים (ספט' 26–דצמ' 27), כל תא נוסחה
  תוכנית עבודה — טבלת היעדים לפי רבעון, בדיוק כמו בדף, פעמיים

הרצה:  python3 src/work_plan_2027.py   → out/Fizikal_תוכנית_עבודה_2027.xlsx
הקובץ נשמר בלי ערכים מחושבים (fullCalcOnLoad) — אקסל/Google Sheets מחשבים בפתיחה.
אימות: הפונקציה model() כאן נותנת את אותם מספרים כמו הנוסחאות (נבדק עם ספריית formulas).
"""
# Port of plan2027.html model() for verification/calibration
BASE_REC=1_000_000; BASE_PART=53_000; MOVE_PRICE=182
TODAY=dict(sales=7660, churnS=740, churnM=2240, up=200)
ONE=[("hw",62.9),("chips",27.4),("dev",50.4),("train",3.2),("tech",3.1),("sms",50)]
PRICES=dict(starter=189,grow=399,scale=549,c1=690,c2=1190)   # preset b (Chen's edits)
MIX=dict(starter=40,grow=30,scale=15,c1=10,c2=5); ADDON=66
COMPS=[49,69,69,49,0,39]; ADOPT=0.64
def arpa():
    comp=sum(COMPS)*ADOPT
    eff={k:PRICES[k]+(comp if k in("c1","c2") else 0) for k in PRICES}
    return sum(eff[k]*MIX[k] for k in MIX)/sum(MIX.values())+ADDON
ARPA=round(arpa())
def phase(i): return 0 if i<4 else i-3 if i<7 else 4 if i<10 else 5 if i<13 else 6
def qk(i): return 0 if i<4 else 1 if i<7 else 2 if i<10 else 3 if i<13 else 4
S={
 "15":dict(clients=[10,12,14,17,27,30,32],meet=[0.48,0.50,0.51,0.52,0.54,0.55,0.55],leads=[82,120,130,140,160,160,160],
   move=[[8,4],[9,4],[9,3],[10,3],[10,2]],churn=[[740,2240],[750,2150],[750,2000]],up=[[10,100],[20,100],[25,120]],
   one=dict(hw=62.9,chips=31,dev=56,train=3.2,tech=3.1,sms=50),part=55000),
 "24":dict(clients=[10,12,15,18,30,36,40],meet=[0.48,0.50,0.51,0.52,0.55,0.58,0.58],leads=[82,180,200,220,250,250,250],
   move=[[8,4],[10,5],[10,3],[10,2],[10,2]],churn=[[740,2240],[750,2000],[750,1750]],up=[[20,100],[30,110],[45,150]],
   one=dict(hw=62.9,chips=36,dev=70,train=3.2,tech=3.1,sms=50),part=61000),
 "30":dict(clients=[10,12,16,20,38,46,54],meet=[0.48,0.50,0.52,0.54,0.58,0.60,0.60],leads=[82,180,220,250,300,300,300],
   move=[[8,4],[10,5],[12,3],[12,2],[12,2]],churn=[[740,2240],[700,1900],[650,1500]],up=[[20,100],[40,120],[60,150]],
   one=dict(hw=70,chips=40,dev=80,train=3.2,tech=3.1,sms=50),part=65000),
}
def model(p):
    R=BASE_REC; mv=0; rows=[]
    for i in range(16):
        ph=phase(i); q=qk(i)
        meet=p["meet"][ph]; leads=p["leads"][ph]; meetings=leads*meet
        extra=max(0,leads*(meet-0.48)*0.77*0.58*0.55)
        n=p["clients"][ph]+(0 if i<4 else extra)
        sales=TODAY["sales"] if i<4 else n*ARPA
        s,c=p["move"][q]; mvNet=(s-c)*MOVE_PRICE
        ch=p["churn"][0 if i<4 else 1 if i<7 else 2]; churn=sum(ch)
        if i==0: upD,upA,up=0,0,TODAY["up"]
        else:
            u=p["up"][0 if i<4 else 1 if i<10 else 2]; upD,upA=u; up=upD*upA
        R+=sales-churn+up; mv+=mvNet
        f=(i+1)/16
        oneBy={k:(b+(p["one"][k]-b)*f)*1000 for k,b in ONE}; one=sum(oneBy.values())
        part=BASE_PART+(p["part"]-BASE_PART)*f
        rows.append(dict(i=i,ph=ph,n=n,meet=meet,meetings=meetings,extra=extra,mv=(s,c),mvNet=mvNet,sales=sales,churn=churn,up=up,upD=upD,upA=upA,R=R,mvCum=mv,one=one,oneBy=oneBy,part=part,tot=R+mv+one+part))
    return rows

import openpyxl, json
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.comments import Comment

S["30"]=dict(clients=[10,12,16,20,34,42,48],meet=[0.48,0.50,0.52,0.54,0.58,0.60,0.60],leads=[82,180,220,250,300,300,300],
   move=[[8,4],[10,5],[12,3],[12,2],[12,2]],churn=[[740,2240],[700,1900],[650,1600]],up=[[20,100],[35,120],[50,150]],
   one=dict(hw=70,chips=40,dev=80,train=3.2,tech=3.1,sms=50),part=65000)

F="Arial"
def font(bold=False,color="000000",size=10): return Font(name=F,bold=bold,color=color,size=size)
BLUE="0000FF"; GREEN="008000"
HFILL=PatternFill("solid",fgColor="1F2A44"); TFILL=PatternFill("solid",fgColor="E8EEF7"); INFILL=PatternFill("solid",fgColor="FFF9DB")
thin=Side(style="thin",color="C9D1E0"); BORDER=Border(top=thin,bottom=thin,left=thin,right=thin)
MONEY='[>=1000000]"₪"0.00,,"M";[>=1000]"₪"#,##0.0,"K";"₪"#,##0'
ILS='"₪"#,##0'
wb=openpyxl.Workbook()
wsP=wb.active; wsP.title="תוכנית עבודה"
wsA=wb.create_sheet("הנחות")
for ws in (wsP,wsA): ws.sheet_view.rightToLeft=True

# ---------------- הנחות ----------------
A={}  # name -> absolute ref on הנחות
def put(ws,r,c,v,name=None,inp=False,fmt=None,note=None):
    cell=ws.cell(row=r,column=c,value=v)
    cell.font=font(color=BLUE if inp else "000000"); cell.border=BORDER
    if inp: cell.fill=INFILL
    if fmt: cell.number_format=fmt
    if note: cell.comment=Comment(note,"Fizikal")
    if name: A[name]=f"'הנחות'!${L(c)}${r}"
    return cell
def label(ws,r,c,t,bold=False):
    cell=ws.cell(row=r,column=c,value=t); cell.font=font(bold=bold); cell.border=BORDER
def header(ws,r,c,t,width=None):
    cell=ws.cell(row=r,column=c,value=t); cell.font=Font(name=F,bold=True,color="FFFFFF"); cell.fill=HFILL
    cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=BORDER

wsA["A1"]="הנחות התוכנית 2027"; wsA["A1"].font=font(True,size=14)
wsA["A2"]="תאים כחולים על רקע צהוב = קלט שאפשר לשנות. שחור = נוסחה. הגיליונות 'מודל' ו'תוכנית עבודה' מחושבים מכאן."; wsA["A2"].font=font(color="555555")
r=4; label(wsA,r,1,"קבועים",True); r+=1
consts=[("avg26","ממוצע מחזור חודשי 2026 (₪)",1261000,"ינואר-אוגוסט 2026, דוח הנה\"ח"),
        ("baseRec","הכנסה חוזרת היום: ריטיינר + מודולים (₪)",1000000,"מעוגל מ-₪1,013K באוגוסט 2026"),
        ("basePart","שיתופי פעולה היום (₪ בחודש)",53000,"מכבי 24 · עמלות 13 · HOWAZIT 9 · UPGRADE 7"),
        ("movePrice","מחיר לקוח MOVE (₪ בחודש)",182,None),
        ("salesToday","מכירות חדשות היום (₪ בחודש, בלי MOVE)",7660,"ממוצע ינואר-יולי 2026 מקובץ המכירות"),
        ("churnS","נטישה היום: קטנים (₪ בחודש)",740,"קובץ נטישות ינואר-יולי 2026, בלי MOVE"),
        ("churnM","נטישה היום: בינוניים (₪ בחודש)",2240,None),
        ("upToday","אפסייל נטו היום (₪ בחודש)",200,None),
        ("meetBase","תיאום פגישה היום (מלידים רלוונטיים)",0.48,"משפך 3 חודשים: 118 פגישות מ-247 לידים רלוונטיים"),
        ("show","הגעה לפגישה",0.77,"91 מ-118"),("close","סגירה מפגישה",0.58,"53 מ-91"),("pay","חלק הסגירות שהן לקוח משלם (לא MOVE)",0.55,"כ-55% מהסגירות אינן MOVE")]
for k,t,v,n in consts:
    label(wsA,r,1,t); put(wsA,r,2,v,k,True,"0%" if v<1 else "#,##0",n); r+=1
r+=1; label(wsA,r,1,"חד-פעמי היום (₪K בחודש, ממוצע 2026)",True); r+=1
ONE_N={"hw":"חומרה: קוראים, QR, בקרי טיבו","chips":"צ'יפים וכרטיסיות","dev":"פיתוח","train":"הדרכות והטמעות","tech":"זמן טכנאי ושליחויות","sms":"SMS"}
for k,b in ONE:
    label(wsA,r,1,ONE_N[k]); put(wsA,r,2,b,"oneB_"+k,True,"0.0","SMS: ₪50K לפי החלטת Chen" if k=="sms" else None); r+=1
r+=1; label(wsA,r,1,"מחירון ללקוחות חדשים",True); r+=1
for c,t in enumerate(["חבילה","מחיר (₪)","תמהיל (%)","רכיבים למותאם (₪)","מחיר אפקטיבי (₪)"],1): header(wsA,r,c,t)
r+=1; pr0=r
PK=[("starter","Starter · עד 50 מתאמנים"),("grow","Grow · 51–120"),("scale","Scale · 121–300"),("c1","מותאם 1 · 301–800"),("c2","מותאם 2 · 800+")]
for k,t in PK:
    label(wsA,r,1,t); put(wsA,r,2,PRICES[k],None,True,"#,##0"); put(wsA,r,3,MIX[k]/100,None,True,"0%")
    if k in("c1","c2"): wsA.cell(row=r,column=4,value="=compAvg").font=font(); wsA.cell(row=r,column=4).number_format="#,##0"
    else: put(wsA,r,4,0,None,False,"#,##0")
    wsA.cell(row=r,column=4).border=BORDER
    c5=wsA.cell(row=r,column=5,value=f"=B{r}+D{r}"); c5.font=font(); c5.number_format="#,##0"; c5.border=BORDER
    r+=1
pr1=r-1
label(wsA,r,1,"תוספות ממוצעות ללקוח (₪)"); put(wsA,r,2,ADDON,"addon",True,"#,##0"); r+=1
label(wsA,r,1,"ממוצע ללקוח חדש (₪) — ARPA",True); c=wsA.cell(row=r,column=2,value=f"=ROUND(SUMPRODUCT(E{pr0}:E{pr1},C{pr0}:C{pr1})/SUM(C{pr0}:C{pr1})+B{r-1},0)"); c.font=font(True); c.number_format="#,##0"; c.border=BORDER
A["arpa"]=f"'הנחות'!$B${r}"; r+=2
label(wsA,r,1,"רכיבים שמתומחרים רק במותאם (₪ בחודש)",True); r+=1; cp0=r
for n,p in zip(["טפסים דיגיטליים","רכישה אינטרנטית","אפליקציה גנרית","ממשק אנרג'ים","סליקה","מיני סייט"],COMPS):
    label(wsA,r,1,n); put(wsA,r,2,p,None,True,"#,##0"); r+=1
cp1=r-1
label(wsA,r,1,"שיעור אימוץ הרכיבים במותאם"); put(wsA,r,2,ADOPT,"adopt",True,"0%"); r+=1
label(wsA,r,1,"תוספת רכיבים ממוצעת למותאם (₪)"); c=wsA.cell(row=r,column=2,value=f"=SUM(B{cp0}:B{cp1})*B{r-1}"); c.font=font(); c.number_format="#,##0"; c.border=BORDER
wb.defined_names.add(openpyxl.workbook.defined_name.DefinedName("compAvg",attr_text=f"'הנחות'!$B${r}")); r+=2

# scenario blocks side by side: 15% at cols A-D, 30% at cols F-I
SC=[("15","צמיחה 15%",1),("30","צמיחה 30%",6)]
PH_N=["רבעון 4 26 (היום)","ינואר 27","פברואר 27","מרץ 27","רבעון 2 27","רבעון 3 27","רבעון 4 27"]
Q_N=["רבעון 4 26","רבעון 1 27","רבעון 2 27","רבעון 3 27","רבעון 4 27"]
P_N=["רבעון 4 26","רבעון 1 27","רבעון 2 27 ואילך"]
U_N=["רבעון 4 26","מחצית 1 2027","מחצית 2 2027"]
R={}  # scenario -> dict of ranges
r0=r
for key,title,c0 in SC:
    p=S[key]; rr=r0; d={}
    label(wsA,rr,c0,title,True); wsA.cell(row=rr,column=c0).font=font(True,size=12); rr+=1
    for i,t in enumerate(["שלב","לקוחות משלמים חדשים בחודש","תיאום פגישה","לידים רלוונטיים בחודש"]): header(wsA,rr,c0+i,t)
    rr+=1; a=rr
    for i,n in enumerate(PH_N):
        label(wsA,rr,c0,n); put(wsA,rr,c0+1,p["clients"][i],None,True,"0"); put(wsA,rr,c0+2,p["meet"][i],None,True,"0%"); put(wsA,rr,c0+3,p["leads"][i],None,True,"0"); rr+=1
    d["clients"]=f"'הנחות'!${L(c0+1)}${a}:${L(c0+1)}${rr-1}"; d["meet"]=f"'הנחות'!${L(c0+2)}${a}:${L(c0+2)}${rr-1}"; d["leads"]=f"'הנחות'!${L(c0+3)}${a}:${L(c0+3)}${rr-1}"
    rr+=1
    for i,t in enumerate(["רבעון","MOVE חתימות בחודש","MOVE ביטולים בחודש"]): header(wsA,rr,c0+i,t)
    rr+=1; a=rr
    for i,n in enumerate(Q_N):
        label(wsA,rr,c0,n); put(wsA,rr,c0+1,p["move"][i][0],None,True,"0"); put(wsA,rr,c0+2,p["move"][i][1],None,True,"0"); rr+=1
    d["mvS"]=f"'הנחות'!${L(c0+1)}${a}:${L(c0+1)}${rr-1}"; d["mvC"]=f"'הנחות'!${L(c0+2)}${a}:${L(c0+2)}${rr-1}"; rr+=1
    for i,t in enumerate(["תקופה","תקרת נטישה קטנים (₪)","תקרת נטישה בינוניים (₪)"]): header(wsA,rr,c0+i,t)
    rr+=1; a=rr
    for i,n in enumerate(P_N):
        label(wsA,rr,c0,n); put(wsA,rr,c0+1,p["churn"][i][0],None,True,"#,##0"); put(wsA,rr,c0+2,p["churn"][i][1],None,True,"#,##0"); rr+=1
    d["chS"]=f"'הנחות'!${L(c0+1)}${a}:${L(c0+1)}${rr-1}"; d["chM"]=f"'הנחות'!${L(c0+2)}${a}:${L(c0+2)}${rr-1}"; rr+=1
    for i,t in enumerate(["תקופה","עסקאות אפסייל בחודש","ממוצע לעסקה (₪)"]): header(wsA,rr,c0+i,t)
    rr+=1; a=rr
    for i,n in enumerate(U_N):
        label(wsA,rr,c0,n); put(wsA,rr,c0+1,p["up"][i][0],None,True,"0"); put(wsA,rr,c0+2,p["up"][i][1],None,True,"#,##0"); rr+=1
    d["upD"]=f"'הנחות'!${L(c0+1)}${a}:${L(c0+1)}${rr-1}"; d["upA"]=f"'הנחות'!${L(c0+2)}${a}:${L(c0+2)}${rr-1}"; rr+=1
    for i,t in enumerate(["חד-פעמי בדצמבר 2027 (₪K בחודש)","יעד"]): header(wsA,rr,c0+i,t)
    rr+=1
    for k,b in ONE:
        label(wsA,rr,c0,ONE_N[k]); put(wsA,rr,c0+1,p["one"][k],None,True,"0.0"); d["oneE_"+k]=f"'הנחות'!${L(c0+1)}${rr}"; rr+=1
    label(wsA,rr,c0,"שיתופי פעולה בדצמבר 2027 (₪ בחודש)"); put(wsA,rr,c0+1,p["part"],None,True,"#,##0"); d["partE"]=f"'הנחות'!${L(c0+1)}${rr}"; rr+=1
    label(wsA,rr,c0,"יעד מחזור דצמבר 2027 (₪)"); put(wsA,rr,c0+1,{"15":1450000,"30":1640000}[key],None,True,"#,##0"); d["target"]=f"'הנחות'!${L(c0+1)}${rr}"
    R[key]=d
for c,w in zip("ABCDEFGHI",[44,16,16,16,16,44,16,16,16]): wsA.column_dimensions[c].width=w

# ---------------- model sheets ----------------
MONTHS=["ספט׳ 26","אוק׳ 26","נוב׳ 26","דצמ׳ 26","ינו׳ 27","פבר׳ 27","מרץ 27","אפר׳ 27","מאי 27","יוני 27","יולי 27","אוג׳ 27","ספט׳ 27","אוק׳ 27","נוב׳ 27","דצמ׳ 27"]
COLS=["חודש","#","שלב","רבעון","תקופת נטישה","תקופת אפסייל","f","לידים","תיאום","פגישות","לקוחות נוספים מתיאום","לקוחות משלמים חדשים","מכירות חדשות (₪)","MOVE חתימות","MOVE ביטולים","MOVE נטו (₪)","MOVE מצטבר (₪)","נטישה (₪)","אפסייל עסקאות","אפסייל ממוצע (₪)","אפסייל (₪)","הכנסה חוזרת (₪)"]+[ONE_N[k] for k,_ in ONE]+["סה\"כ חד-פעמי (₪)","שיתופי פעולה (₪)","מחזור (₪)"]
MC={n:i+1 for i,n in enumerate(COLS)}
ONE_TOT="סה\"כ חד-פעמי (₪)"
def mc(n): return L(MC[n])
MROW={}
for key,title,_ in SC:
    ws=wb.create_sheet(f"מודל {key}"); ws.sheet_view.rightToLeft=True; ws.freeze_panes="B3"
    ws["A1"]=f"מודל חודשי — {title}. כל התאים נוסחאות; הקלט בגיליון 'הנחות'."; ws["A1"].font=font(True,size=12)
    for i,n in enumerate(COLS,1): header(ws,2,i,n); ws.column_dimensions[L(i)].width=13
    ws.column_dimensions["A"].width=10; ws.row_dimensions[2].height=42
    d=R[key]; rows=[]
    for i in range(16):
        r=3+i; rows.append(r); prev=r-1
        ph=phase(i); q=qk(i); cp=0 if i<4 else 1 if i<7 else 2; up=0 if i<4 else 1 if i<10 else 2
        vals={"חודש":MONTHS[i],"#":i,"שלב":ph,"רבעון":q,"תקופת נטישה":cp,"תקופת אפסייל":up,
          "f":f"=({mc('#')}{r}+1)/16",
          "לידים":f"=INDEX({d['leads']},{mc('שלב')}{r}+1)","תיאום":f"=INDEX({d['meet']},{mc('שלב')}{r}+1)",
          "פגישות":f"={mc('לידים')}{r}*{mc('תיאום')}{r}",
          "לקוחות נוספים מתיאום":f"=MAX(0,{mc('לידים')}{r}*({mc('תיאום')}{r}-{A['meetBase']})*{A['show']}*{A['close']}*{A['pay']})",
          "לקוחות משלמים חדשים":f"=INDEX({d['clients']},{mc('שלב')}{r}+1)+IF({mc('#')}{r}<4,0,{mc('לקוחות נוספים מתיאום')}{r})",
          "מכירות חדשות (₪)":f"=IF({mc('#')}{r}<4,{A['salesToday']},{mc('לקוחות משלמים חדשים')}{r}*{A['arpa']})",
          "MOVE חתימות":f"=INDEX({d['mvS']},{mc('רבעון')}{r}+1)","MOVE ביטולים":f"=INDEX({d['mvC']},{mc('רבעון')}{r}+1)",
          "MOVE נטו (₪)":f"=({mc('MOVE חתימות')}{r}-{mc('MOVE ביטולים')}{r})*{A['movePrice']}",
          "MOVE מצטבר (₪)":(f"={mc('MOVE נטו (₪)')}{r}" if i==0 else f"={mc('MOVE מצטבר (₪)')}{prev}+{mc('MOVE נטו (₪)')}{r}"),
          "נטישה (₪)":f"=INDEX({d['chS']},{mc('תקופת נטישה')}{r}+1)+INDEX({d['chM']},{mc('תקופת נטישה')}{r}+1)",
          "אפסייל עסקאות":(f"=0" if i==0 else f"=INDEX({d['upD']},{mc('תקופת אפסייל')}{r}+1)"),
          "אפסייל ממוצע (₪)":(f"=0" if i==0 else f"=INDEX({d['upA']},{mc('תקופת אפסייל')}{r}+1)"),
          "אפסייל (₪)":(f"={A['upToday']}" if i==0 else f"={mc('אפסייל עסקאות')}{r}*{mc('אפסייל ממוצע (₪)')}{r}"),
          "הכנסה חוזרת (₪)":f"={(A['baseRec'] if i==0 else mc('הכנסה חוזרת (₪)')+str(prev))}+{mc('מכירות חדשות (₪)')}{r}-{mc('נטישה (₪)')}{r}+{mc('אפסייל (₪)')}{r}",
          "שיתופי פעולה (₪)":f"={A['basePart']}+({d['partE']}-{A['basePart']})*{mc('f')}{r}"}
        for k,_ in ONE: vals[ONE_N[k]]=f"=({A['oneB_'+k]}+({d['oneE_'+k]}-{A['oneB_'+k]})*{mc('f')}{r})*1000"
        vals[ONE_TOT]=f"=SUM({mc(ONE_N['hw'])}{r}:{mc(ONE_N['sms'])}{r})"
        vals["מחזור (₪)"]=f"={mc('הכנסה חוזרת (₪)')}{r}+{mc('MOVE מצטבר (₪)')}{r}+{mc(ONE_TOT)}{r}+{mc('שיתופי פעולה (₪)')}{r}"
        for n in COLS:
            c=ws.cell(row=r,column=MC[n],value=vals[n]); c.border=BORDER
            c.font=font(color=(BLUE if n in("#","שלב","רבעון","תקופת נטישה","תקופת אפסייל") else "000000"))
            if n=="תיאום": c.number_format="0%"
            elif n=="f": c.number_format="0.000"
            elif "₪" in n: c.number_format="#,##0"
            elif n in("פגישות","לקוחות נוספים מתיאום","לקוחות משלמים חדשים"): c.number_format="0.0"
    ws.cell(row=2,column=MC["שלב"]).comment=Comment("אינדקס מבני: 0 = רבעון 4 26, 1-3 = ינואר-מרץ 27, 4-6 = רבעונים 2-4 של 2027. מצביע על טבלת השלבים בגיליון 'הנחות'.","Fizikal")
    MROW[key]=rows

# ---------------- תוכנית עבודה ----------------
ws=wsP
ws["A1"]="Fizikal · תוכנית עבודה 2027 — היעדים לפי רבעון"; ws["A1"].font=font(True,size=15)
ws["A2"]="מנופים חוזרים: הקצב בכל חודש ברבעון · חד-פעמי: ממוצע חודשי · מחזור: החודש האחרון ברבעון. כל התאים נוסחאות מהגיליונות 'מודל' ו'הנחות'."; ws["A2"].font=font(color="555555")
ws.column_dimensions["A"].width=46
for c in "BCDEFG": ws.column_dimensions[c].width=20
ws.column_dimensions["D"].width=34
QCOLS=["היום"]+Q_N
QMON=[[1,2,3],[4,5,6],[7,8,9],[10,11,12],[13,14,15]]
def tbl(top,key,title):
    rows=MROW[key]; mk=f"'מודל {key}'!"
    def ref(col,i): return f"{mk}{mc(col)}{rows[i]}"
    c=ws.cell(row=top,column=1,value=title); c.font=font(True,size=13)
    g=ws.cell(row=top,column=2,value=f"=\"דצמבר 2027: \"&TEXT({ref('מחזור (₪)',15)}/1000000,\"0.00\")&\"M ₪ · \"&TEXT({ref('מחזור (₪)',15)}/{A['avg26']}-1,\"+0%\")&\" מול ממוצע 2026 · יעד \"&TEXT({R[key]['target']}/1000000,\"0.00\")&\"M\""); g.font=font(color="555555")
    ws.merge_cells(start_row=top,start_column=2,end_row=top,end_column=7)
    hr=top+1
    header(ws,hr,1,"")
    for i,t in enumerate(QCOLS,2): header(ws,hr,i,t)
    ws.row_dimensions[hr].height=24
    fmtN=lambda x: f"IF({x}=INT({x}),TEXT({x},\"0\"),TEXT({x},\"0.0\"))"
    meetTxt=lambda i: f"TEXT({ref('תיאום',i)},\"0%\")&\" (\"&TEXT({ref('פגישות',i)},\"0\")&\")\""
    kTxt=lambda i: f"\"₪\"&TEXT({ref('מכירות חדשות (₪)',i)}/1000,\"0.0\")&\"K\""
    def chain(f): return "="+'&" ← "&'.join(f(i) for i in (4,5,6))
    lines=[
     ("תיאום פגישה מרלוונטי (פגישות בחודש)", f"=TEXT({A['meetBase']},\"0%\")&\" (\"&TEXT(INDEX({R[key]['leads']},1)*{A['meetBase']},\"0\")&\")\"",
        [("="+meetTxt(ms[0])) if k!=1 else chain(meetTxt) for k,ms in enumerate(QMON)], None),
     ("לקוחות משלמים חדשים בחודש", f"=INDEX({R[key]['clients']},1)",
        [("="+fmtN(ref('לקוחות משלמים חדשים',ms[0]))) if k!=1 else chain(lambda i: fmtN(ref('לקוחות משלמים חדשים',i))) for k,ms in enumerate(QMON)], "0"),
     (f"=\"מכירות חדשות (₪ בחודש, ממוצע ₪\"&TEXT({A['arpa']},\"#,##0\")&\")\"", f"={A['salesToday']}",
        [(f"={ref('מכירות חדשות (₪)',ms[0])}") if k!=1 else chain(kTxt) for k,ms in enumerate(QMON)], MONEY),
     ("MOVE נטו (חתימות − ביטולים)", f"=INDEX({R[key]['mvS']},1)&\"−\"&INDEX({R[key]['mvC']},1)&\" = ₪\"&TEXT((INDEX({R[key]['mvS']},1)-INDEX({R[key]['mvC']},1))*{A['movePrice']},\"#,##0\")",
        [f"={ref('MOVE חתימות',ms[0])}&\"−\"&{ref('MOVE ביטולים',ms[0])}&\" = ₪\"&TEXT({ref('MOVE נטו (₪)',ms[0])},\"#,##0\")" for ms in QMON], None),
     ("נטישה (תקרה בחודש)", f"={A['churnS']}+{A['churnM']}", [f"={ref('נטישה (₪)',ms[0])}" for ms in QMON], MONEY),
     ("אפסייל נטו (עסקאות × ממוצע)", f"={A['upToday']}", [f"={ref('אפסייל עסקאות',ms[0])}&\" × ₪\"&TEXT({ref('אפסייל ממוצע (₪)',ms[0])},\"#,##0\")" for ms in QMON], MONEY),
     ("חד-פעמי כולל SMS (ממוצע)", f"=SUM({A['oneB_hw']},{A['oneB_chips']},{A['oneB_dev']},{A['oneB_train']},{A['oneB_tech']},{A['oneB_sms']})*1000",
        [f"=AVERAGE({mk}{mc(ONE_TOT)}{rows[ms[0]]}:{mc(ONE_TOT)}{rows[ms[2]]})" for ms in QMON], MONEY),
     ("שיתופי פעולה", f"={A['basePart']}", [f"={ref('שיתופי פעולה (₪)',ms[2])}" for ms in QMON], MONEY),
     ("מחזור בחודש האחרון ברבעון", f"={A['baseRec']}+SUM({A['oneB_hw']},{A['oneB_chips']},{A['oneB_dev']},{A['oneB_train']},{A['oneB_tech']},{A['oneB_sms']})*1000+{A['basePart']}",
        [f"={ref('מחזור (₪)',ms[2])}" for ms in QMON], MONEY),
    ]
    for j,(lab,today,qs,fmt) in enumerate(lines):
        r=hr+1+j; total=(j==len(lines)-1); neg=lab.startswith("נטישה")
        c=ws.cell(row=r,column=1,value=lab); c.font=font(bold=total); c.border=BORDER
        for ci,v in enumerate([today]+qs,2):
            c=ws.cell(row=r,column=ci,value=v); c.border=BORDER
            c.font=font(bold=total,color=("C62828" if (neg and ci>2) else "000000"))
            c.alignment=Alignment(horizontal="center")
            if fmt: c.number_format=fmt
        if total:
            for ci in range(1,8): ws.cell(row=r,column=ci).fill=TFILL
    nr=hr+1+len(lines)
    ws.cell(row=nr,column=1,value={"15":"במרץ 2027: 17+ לקוחות משלמים חדשים בחודש = התוכנית בטוחה.","30":"במרץ 2027: 20+ לקוחות משלמים חדשים = התוכנית בטוחה. דורש הטמעה עצמית לקטנים ברבעון 2, אחרת 3 אנשי CS לא עומדים בקצב."}[key]).font=font(color="555555")
    return nr+2
nxt=tbl(4,"15","צמיחה 15% — דצמבר 2027 ₪1.45M")
tbl(nxt,"30","צמיחה 30% — דצמבר 2027 ₪1.64M")
wb.calculation.fullCalcOnLoad=True
out="/home/user/Claude-projects/fizikal-crm/out/Fizikal_תוכנית_עבודה_2027.xlsx"
wb.save(out); print("saved",out)
