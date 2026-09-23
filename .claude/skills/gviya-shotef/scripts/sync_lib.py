#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Self-contained report reader for the weekly aging XLSX (replaces the sync.py dependency that
lived in Chen's zip and does not survive container resets). Header lookup by TEXT only."""
import openpyxl, datetime, re, json
MONTHS = ["ינואר","פברואר","מרץ","אפריל","מאי","יוני","יולי","אוגוסט","ספטמבר","אוקטובר","נובמבר","דצמבר"]
def digits(v): return re.sub(r"\D", "", str(v or ""))
def report_date(sheet):
    m = re.match(r"(\d{1,2})\.(\d{1,2})\.(\d{2,4})", sheet.strip())
    if not m: return datetime.date.today().isoformat()
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    return datetime.date(y + 2000 if y < 100 else y, mo, d).isoformat()
def read_report(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    name = sheet or wb.sheetnames[-1]; ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    hdr = next(i for i, r in enumerate(rows[:8]) if any(str(c).strip() == "שם לקוח" for c in r if c))
    cols = [str(c).strip() if c is not None else "" for c in rows[hdr]]
    ci = {}
    for j, n in enumerate(cols):
        if n and n not in ci: ci[n] = j          # first-wins: duplicate 'יתרה לתשלום' trap
    bal_j = ci.get("יתרה לתשלום")
    if bal_j is None: raise SystemExit("no 'יתרה לתשלום' column in sheet %s" % name)
    aging = []
    for j, raw in enumerate(rows[hdr]):
        n = cols[j]
        if isinstance(raw, datetime.datetime): aging.append((j, raw.date(), "month"))
        elif n in MONTHS: aging.append((j, n, "lastyear"))
        elif n == "שנים קודמות": aging.append((j, n, "ancient"))
    out = {}
    for r in rows[hdr + 1:]:
        code = r[ci["קוד לקוח"]] if "קוד לקוח" in ci else None
        if code is None or not str(code).strip().isdigit(): continue
        def num(x):
            try: return float(x) if x not in (None, "") else 0.0
            except (TypeError, ValueError): return 0.0
        buckets = [(l, k, round(num(r[j]))) for j, l, k in aging if num(r[j])]
        out[str(code).strip()] = {"code": str(code).strip(), "name": str(r[ci["שם לקוח"]] or "").strip(),
            "hp": digits(r[ci["ח.פ."]]) if "ח.פ." in ci else "",
            "terms": str(r[ci["שוטף"]] or "").strip() if "שוטף" in ci else "",
            "bal": round(num(r[bal_j])), "buckets": buckets}
    return name, out, report_date(name)
def debt_age(rec, asof):
    if not rec["buckets"]: return None, None
    today = datetime.date.fromisoformat(asof)
    if any(k == "ancient" for l, k, v in rec["buckets"]):
        return "שנים קודמות", (today - datetime.date(today.year - 1, 1, 1)).days
    olds = [l for l, k, v in rec["buckets"] if k == "lastyear"]
    if olds:
        mi = min(MONTHS.index(o) for o in olds)
        return "%s (שנה קודמת)" % MONTHS[mi], (today - datetime.date(today.year - 1, mi + 1, 1)).days
    ms = [l for l, k, v in rec["buckets"] if k == "month"]
    if ms:
        o = min(ms); return o.strftime("%m/%y"), (today - o).days
    return None, None
def overdue_of(rec, asof):
    cur = datetime.date.fromisoformat(asof).replace(day=1)
    return round(sum(v for l, k, v in rec["buckets"] if v > 0 and (k != "month" or l < cur)))
def build_snapshot(rep, sheet, asof):
    return {"date": asof, "sheet": sheet, "clients": {c: [r["bal"], overdue_of(r, asof)] for c, r in rep.items()}}
def board_from_dump(raw_path):
    """monday get_board_items_page dump → {code: {...}} keyed by קוד לקוח."""
    b = json.load(open(raw_path)); out = {}
    for it in b["items"]:
        cv = it["column_values"]; code = (cv.get("text_mm69gg7g") or "").strip()
        if not code: continue
        out[code] = {"id": it["id"], "name": it["name"], "stage": cv.get("color_mm69ws8c"), "debt": cv.get("numeric_mm69dbtr"),
                     "bal": cv.get("color_mm69xcy2"), "terms": cv.get("text_mm69ew8e"), "age": cv.get("text_mm6970hn"),
                     "report": cv.get("date_mm69pvs6"), "promised": cv.get("date_mm7170br") or None}
    return out
